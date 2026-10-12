from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(slide_layout)

def add_text_box(slide, left, top, width, height, text, font_size=14, bold=False, color=(0,0,0)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    p.alignment = PP_ALIGN.CENTER
    return txBox

def add_rectangle(slide, left, top, width, height, fill_color, text='', font_size=12, text_color=(255,255,255)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill_color)
    shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = RGBColor(*text_color)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape

def add_line(slide, start_x, start_y, end_x, end_y, color=(100,100,100)):
    from pptx.enum.shapes import MSO_CONNECTOR
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start_x), Inches(start_y),
        Inches(end_x), Inches(end_y)
    )
    connector.line.color.rgb = RGBColor(*color)
    connector.line.width = Pt(2)
    return connector

title_box = add_rectangle(slide, 0.3, 0.2, 12.7, 0.7, (47, 84, 150), 'STAEformer 嵌入层结构（Embedding Layer）', 26, (255, 255, 255))

add_text_box(slide, 0.3, 1.0, 6, 0.4, '输入：x [batch_size, in_steps=12, num_nodes, input_dim=3]', 16, True, (0, 0, 0))

input_box = add_rectangle(slide, 0.5, 1.5, 2.2, 0.6, (200, 200, 200), 'x: (12, 207, 3)', 14, (0, 0, 0))

y_pos = 2.4
gap = 1.8

add_rectangle(slide, 0.5, y_pos, 1.5, 0.6, (255, 192, 0), 'Input Proj', 13, (255, 255, 255))
add_rectangle(slide, 0.5, y_pos+0.7, 1.5, 0.5, (255, 192, 0), 'Linear(3→24)', 11, (255, 255, 255))

add_rectangle(slide, 0.5+gap, y_pos, 1.5, 0.6, (112, 173, 71), 'TOD Embed', 13, (255, 255, 255))
add_rectangle(slide, 0.5+gap, y_pos+0.7, 1.5, 0.5, (112, 173, 71), 'Embed(288→24)', 11, (255, 255, 255))

add_rectangle(slide, 0.5+gap*2, y_pos, 1.5, 0.6, (112, 60, 160), 'DOW Embed', 13, (255, 255, 255))
add_rectangle(slide, 0.5+gap*2, y_pos+0.7, 1.5, 0.5, (112, 60, 160), 'Embed(7→24)', 11, (255, 255, 255))

add_rectangle(slide, 0.5+gap*3, y_pos, 1.8, 0.6, (197, 90, 17), 'Adaptive Embed', 13, (255, 255, 255))
add_rectangle(slide, 0.5+gap*3, y_pos+0.7, 1.8, 0.5, (197, 90, 17), 'Param(12×207×80)', 11, (255, 255, 255))

add_line(slide, 1.6, 2.1, 1.0, y_pos-0.05)
add_line(slide, 1.6, 2.1, 2.8, y_pos-0.05)
add_line(slide, 1.6, 2.1, 4.6, y_pos-0.05)
add_line(slide, 1.6, 2.1, 6.4, y_pos-0.05)
add_line(slide, 1.0, y_pos-0.05, 6.4, y_pos-0.05)

add_line(slide, 1.0, y_pos+0.15, 1.0, y_pos+0.55)
add_line(slide, 2.8, y_pos+0.15, 2.8, y_pos+0.55)
add_line(slide, 4.6, y_pos+0.15, 4.6, y_pos+0.55)
add_line(slide, 6.4, y_pos+0.15, 6.4, y_pos+0.55)

add_text_box(slide, 0.5, 3.5, 1.5, 0.3, '24维', 11, False, (80, 80, 80))
add_text_box(slide, 0.5+gap, 3.5, 1.5, 0.3, '24维', 11, False, (80, 80, 80))
add_text_box(slide, 0.5+gap*2, 3.5, 1.5, 0.3, '24维', 11, False, (80, 80, 80))
add_text_box(slide, 0.5+gap*3, 3.5, 1.8, 0.3, '80维', 11, False, (80, 80, 80))

add_line(slide, 1.0, 3.8, 1.0, 4.3)
add_line(slide, 2.8, 3.8, 2.8, 4.3)
add_line(slide, 4.6, 3.8, 4.6, 4.3)
add_line(slide, 6.4, 3.8, 6.4, 4.3)

add_line(slide, 1.0, 4.3, 1.0, 4.6)
add_line(slide, 2.8, 4.3, 2.8, 4.6)
add_line(slide, 4.6, 4.3, 4.6, 4.6)
add_line(slide, 6.4, 4.3, 6.4, 4.6)

add_line(slide, 1.0, 4.6, 2.8, 4.6)
add_line(slide, 2.8, 4.6, 4.6, 4.6)
add_line(slide, 4.6, 4.6, 6.4, 4.6)
add_line(slide, 6.4, 4.6, 7.3, 4.6)

for i in range(4):
    bx = 0.5 + i * gap if i < 3 else 0.5 + 3 * gap
    add_text_box(slide, bx, 4.3, 0.3, 0.3, '│', 14, False, (80, 80, 80))
    add_text_box(slide, bx, 4.45, 0.3, 0.3, '│', 14, False, (80, 80, 80))
    add_text_box(slide, bx, 4.6, 0.3, 0.3, '│', 14, False, (80, 80, 80))

add_text_box(slide, 0.3, 4.3, 0.6, 0.3, '│', 14, False, (80, 80, 80))
add_text_box(slide, 0.3, 4.45, 0.6, 0.3, '│', 14, False, (80, 80, 80))
add_text_box(slide, 0.3, 4.6, 0.6, 0.3, '│', 14, False, (80, 80, 80))

output_box = add_rectangle(slide, 7.3, 4.15, 1.5, 0.6, (47, 84, 150), 'Concat', 14, (255, 255, 255))
output_final = add_rectangle(slide, 7.3, 4.9, 1.5, 0.5, (47, 84, 150), '(12,207,152)', 14, (255, 255, 255))
add_line(slide, 7.55, 4.7, 7.55, 4.9)

right_x = 9.5
add_text_box(slide, right_x, 1.5, 3.3, 0.4, '各嵌入层说明', 18, True, (0, 0, 0))

desc_y = 2.0
gap_y = 1.15

add_rectangle(slide, right_x, desc_y, 3.3, 0.5, (255, 192, 0), '', 12)
add_text_box(slide, right_x+0.1, desc_y+0.05, 3.1, 0.45, 'Input Projection: 24维', 13, True, (0, 0, 0))
add_text_box(slide, right_x, desc_y+0.5, 3.3, 0.6, '将原始流量特征投影到24维空间', 11, False, (80, 80, 80))

add_rectangle(slide, right_x, desc_y+gap_y, 3.3, 0.5, (112, 173, 71), '', 12)
add_text_box(slide, right_x+0.1, desc_y+gap_y+0.05, 3.1, 0.45, 'TOD Embedding: 24维', 13, True, (0, 0, 0))
add_text_box(slide, right_x, desc_y+gap_y+0.5, 3.3, 0.6, '编码一天内的时间位置\n288个时间槽（5分钟间隔）', 11, False, (80, 80, 80))

add_rectangle(slide, right_x, desc_y+gap_y*2, 3.3, 0.5, (112, 60, 160), '', 12)
add_text_box(slide, right_x+0.1, desc_y+gap_y*2+0.05, 3.1, 0.45, 'DOW Embedding: 24维', 13, True, (0, 0, 0))
add_text_box(slide, right_x, desc_y+gap_y*2+0.5, 3.3, 0.6, '编码星期几（1-7）\n区分工作日与周末模式', 11, False, (80, 80, 80))

add_rectangle(slide, right_x, desc_y+gap_y*3, 3.3, 0.5, (197, 90, 17), '', 12)
add_text_box(slide, right_x+0.1, desc_y+gap_y*3+0.05, 3.1, 0.45, 'Adaptive Embedding: 80维', 13, True, (0, 0, 0))
add_text_box(slide, right_x, desc_y+gap_y*3+0.5, 3.3, 0.6, '可学习参数矩阵\n捕获数据特有的时空模式', 11, False, (80, 80, 80))

formula_box = add_rectangle(slide, 0.3, 6.2, 12.7, 0.8, (47, 84, 150), 
    'model_dim = input_embedding(24) + tod_embedding(24) + dow_embedding(24) + adaptive_embedding(80) = 152', 
    16, (255, 255, 255))

prs.save(r'D:\大三下\智能交通\STAEformer_嵌入层介绍.pptx')
print('PPT已保存: D:\\大三下\\智能交通\\STAEformer_嵌入层介绍.pptx')

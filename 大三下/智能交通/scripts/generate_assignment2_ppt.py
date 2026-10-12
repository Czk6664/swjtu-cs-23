from __future__ import annotations

import csv
import math
from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
ASSIGNMENT_DIR = ROOT / "大作业2"
SLICE_DIR = ASSIGNMENT_DIR / "ppt切片"
EXP_DIR = ASSIGNMENT_DIR / "exp_C_704_b8_w0_e50" / "exp_C_704_b8_w0_e50"
REPORT_FIG_DIR = ASSIGNMENT_DIR / "课程设计模版_2026年" / "figures" / "tt100k_experiment"
OUT_DIR = ASSIGNMENT_DIR / "生成汇报PPT"
PAGE_DIR = OUT_DIR / "pages"
PPTX_PATH = OUT_DIR / "大作业2_交通标志检测汇报_生成版.pptx"

W, H = 2048, 1152

FONT_REG = Path("C:/Windows/Fonts/Noto Sans SC (TrueType).otf")
FONT_MED = Path("C:/Windows/Fonts/Noto Sans SC Medium (TrueType).otf")
FONT_BOLD = Path("C:/Windows/Fonts/Noto Sans SC Bold (TrueType).otf")
if not FONT_REG.exists():
    FONT_REG = Path("C:/Windows/Fonts/msyh.ttc")
if not FONT_BOLD.exists():
    FONT_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")
if not FONT_MED.exists():
    FONT_MED = FONT_BOLD


def font(size: int, weight: str = "regular") -> ImageFont.FreeTypeFont:
    path = {"regular": FONT_REG, "medium": FONT_MED, "bold": FONT_BOLD}.get(weight, FONT_REG)
    return ImageFont.truetype(str(path), size=size)


def text_size(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def fit_font(draw: ImageDraw.ImageDraw, text: str, max_width: int, size: int, weight: str = "bold", min_size: int = 28):
    while size > min_size:
        fnt = font(size, weight)
        if text_size(draw, text, fnt)[0] <= max_width:
            return fnt
        size -= 2
    return font(min_size, weight)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
            continue
        current = ""
        for ch in raw:
            candidate = current + ch
            if text_size(draw, candidate, fnt)[0] <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    fnt: ImageFont.FreeTypeFont,
    fill: tuple[int, int, int] | tuple[int, int, int, int],
    max_width: int,
    line_gap: int = 10,
):
    x, y = xy
    for line in wrap_text(draw, text, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += text_size(draw, line or "口", fnt)[1] + line_gap
    return y


def rounded(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], radius: int, fill, outline=None, width: int = 1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def paste_fit(canvas: Image.Image, img_path: Path, box: tuple[int, int, int, int], radius: int = 0, bg=(255, 255, 255)):
    if not img_path.exists():
        return
    x1, y1, x2, y2 = box
    target_w, target_h = x2 - x1, y2 - y1
    src = Image.open(img_path).convert("RGB")
    src.thumbnail((target_w, target_h), Image.Resampling.LANCZOS)
    panel = Image.new("RGB", (target_w, target_h), bg)
    px = (target_w - src.width) // 2
    py = (target_h - src.height) // 2
    panel.paste(src, (px, py))
    if radius:
        mask = Image.new("L", panel.size, 0)
        mdraw = ImageDraw.Draw(mask)
        mdraw.rounded_rectangle((0, 0, target_w, target_h), radius=radius, fill=255)
        canvas.paste(panel, (x1, y1), mask)
    else:
        canvas.paste(panel, (x1, y1))


def slice_image(index: int) -> Path:
    matches = sorted(SLICE_DIR.glob(f"*_{index:02d}.png"))
    if not matches:
        matches = sorted(SLICE_DIR.glob("*.png"))
    return matches[0]


def base_from_slice(index: int, dark: bool = False, blur: int = 0) -> Image.Image:
    img = Image.open(slice_image(index)).convert("RGB").resize((W, H), Image.Resampling.LANCZOS)
    if blur:
        img = img.filter(ImageFilter.GaussianBlur(blur))
    if dark:
        overlay = Image.new("RGB", (W, H), (40, 40, 40))
        img = Image.blend(img, overlay, 0.38)
    return img


def content_bg() -> Image.Image:
    bg = base_from_slice(10, blur=18)
    wash = Image.new("RGB", (W, H), (246, 246, 244))
    return Image.blend(bg, wash, 0.72)


def cover_bg() -> Image.Image:
    return base_from_slice(1, dark=True, blur=1)


def draw_logo_text(draw: ImageDraw.ImageDraw, white: bool = False):
    fill = (250, 250, 250) if white else (70, 70, 70)
    draw.text((W - 285, 80), "SWJTU", font=font(34, "bold"), fill=fill)
    draw.text((W - 405, 123), "Southwest Jiaotong University", font=font(22), fill=fill)


def draw_page_no(draw: ImageDraw.ImageDraw, page: int):
    draw.text((W - 105, H - 88), str(page), font=font(36), fill=(105, 105, 105))


def draw_title_bar(draw: ImageDraw.ImageDraw, title: str, page: int, section: str | None = None):
    draw.rectangle((0, 0, W, 160), fill=(247, 247, 246, 235))
    draw.rectangle((0, 160, W, 168), fill=(86, 86, 84))
    draw.ellipse((82, 68, 116, 102), fill=(88, 88, 86))
    draw.ellipse((122, 68, 156, 102), fill=(145, 145, 142))
    draw.ellipse((162, 68, 196, 102), fill=(205, 205, 200))
    draw.text((225, 58), title, font=fit_font(draw, title, 1200, 62), fill=(45, 45, 45))
    if section:
        rounded(draw, (1510, 45, 1905, 115), 10, fill=(92, 92, 90))
        draw.text((1548, 59), section, font=font(30, "bold"), fill=(255, 255, 255))
    draw_logo_text(draw)
    draw_page_no(draw, page)


def draw_bullets(
    draw: ImageDraw.ImageDraw,
    bullets: Sequence[str],
    x: int,
    y: int,
    max_width: int,
    size: int = 38,
    fill=(255, 255, 255),
    bullet_fill=(255, 255, 255),
    gap: int = 20,
):
    fnt = font(size, "medium")
    for item in bullets:
        draw.ellipse((x, y + 15, x + 12, y + 27), fill=bullet_fill)
        y = draw_wrapped(draw, item, (x + 38, y), fnt, fill, max_width, line_gap=8) + gap
    return y


def draw_light_bullets(draw, bullets, x, y, max_width, size=34, gap=22):
    fnt = font(size, "medium")
    for item in bullets:
        draw.ellipse((x, y + 12, x + 11, y + 23), fill=(70, 70, 68))
        y = draw_wrapped(draw, item, (x + 34, y), fnt, (45, 45, 43), max_width, line_gap=8) + gap
    return y


def slide_cover(page: int) -> Image.Image:
    img = cover_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_logo_text(draw, white=True)
    title = "基于 TT100K 的交通标志检测方法研究与 YOLO 轻量化实现"
    fnt = fit_font(draw, title, 1660, 86, "bold", 56)
    tw, th = text_size(draw, title, fnt)
    draw.text(((W - tw) // 2, 435), title, font=fnt, fill=(255, 255, 255))
    draw.rectangle((290, 555, W - 290, 560), fill=(255, 255, 255, 180))
    info = ["汇报人：程泽凯", "组员：程泽凯、王晨羽、金井直", "2026年5月"]
    y = 720
    for line in info:
        f = font(42 if "汇报" in line else 34, "medium")
        tw, _ = text_size(draw, line, f)
        draw.text(((W - tw) // 2, y), line, font=f, fill=(250, 250, 250))
        y += 74
    draw.text((460, H - 96), "谈 实 场 华    自 强 不 息", font=font(28), fill=(190, 190, 190))
    return img


def slide_toc(page: int, items: Sequence[str]) -> Image.Image:
    img = base_from_slice(2, blur=1)
    draw = ImageDraw.Draw(img, "RGBA")
    draw.rectangle((80, 60, 1320, 1040), fill=(246, 246, 244, 215))
    draw.text((1390, 178), "目 录", font=font(96, "medium"), fill=(42, 42, 42))
    draw.text((1394, 365), "CONTENTS", font=font(54), fill=(42, 42, 42))
    y = 135
    for i, item in enumerate(items, 1):
        draw.ellipse((205, y, 305, y + 100), fill=(70, 70, 70))
        draw.text((232, y + 24), f"{i:02d}", font=font(38, "medium"), fill=(255, 255, 255))
        draw.text((380, y + 14), item, font=font(54, "bold"), fill=(42, 42, 42))
        y += 165
    return img


def slide_section(page: int, number: int, title: str) -> Image.Image:
    img = base_from_slice(3)
    draw = ImageDraw.Draw(img, "RGBA")
    rounded(draw, (455, 250, 1815, 885), 28, fill=(255, 255, 255, 250))
    draw.text((590, 335), f"{number:02d}", font=font(142, "bold"), fill=(92, 92, 90))
    draw.text((590, 555), title, font=fit_font(draw, title, 1040, 74, "bold"), fill=(65, 65, 63))
    draw.line((590, 775, 800, 775), fill=(92, 92, 90), width=5)
    draw_page_no(draw, page)
    return img


def slide_text(page: int, title: str, bullets: Sequence[str], section: str = "", note: str | None = None) -> Image.Image:
    img = content_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_title_bar(draw, title, page, section)
    draw.rectangle((0, 700, W, H), fill=(76, 76, 74, 235))
    y = 235
    for idx, item in enumerate(bullets[:3]):
        x = 150 + idx * 620
        rounded(draw, (x, y, x + 520, 610), 18, fill=(255, 255, 255, 232))
        draw.text((x + 38, y + 38), f"{idx + 1:02d}", font=font(54, "bold"), fill=(90, 90, 88))
        draw_wrapped(draw, item, (x + 38, y + 130), font(34, "medium"), (45, 45, 43), 450, line_gap=10)
    if len(bullets) > 3:
        draw_bullets(draw, bullets[3:], 160, 755, 1720, size=36)
    if note:
        rounded(draw, (1410, 778, 1895, 1018), 16, fill=(255, 255, 255, 32), outline=(235, 235, 235, 80), width=2)
        draw_wrapped(draw, note, (1450, 820), font(30, "medium"), (255, 255, 255), 410, line_gap=8)
    return img


def slide_two_col(page: int, title: str, left_title: str, left_items: Sequence[str], right_title: str, right_items: Sequence[str], section: str = ""):
    img = content_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_title_bar(draw, title, page, section)
    rounded(draw, (110, 230, 980, 1010), 18, fill=(255, 255, 255, 235))
    rounded(draw, (1070, 230, 1940, 1010), 18, fill=(76, 76, 74, 235))
    draw.text((165, 285), left_title, font=font(48, "bold"), fill=(45, 45, 43))
    draw_light_bullets(draw, left_items, 170, 380, 735, size=33)
    draw.text((1128, 285), right_title, font=font(48, "bold"), fill=(255, 255, 255))
    draw_bullets(draw, right_items, 1135, 380, 735, size=33)
    return img


def slide_table(page: int, title: str, headers: Sequence[str], rows: Sequence[Sequence[str]], section: str = ""):
    img = content_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_title_bar(draw, title, page, section)
    x, y = 130, 250
    table_w = W - 260
    row_h = 86
    col_w = [int(table_w * r) for r in ([0.34, 0.17, 0.17, 0.16, 0.16] if len(headers) == 5 else [0.27, 0.27, 0.23, 0.23])]
    rounded(draw, (x, y, x + table_w, y + row_h * (len(rows) + 1)), 16, fill=(255, 255, 255, 240))
    cx = x
    draw.rectangle((x, y, x + table_w, y + row_h), fill=(70, 70, 68))
    for i, h in enumerate(headers):
        draw.text((cx + 24, y + 24), h, font=font(28, "bold"), fill=(255, 255, 255))
        cx += col_w[i]
    for r, row in enumerate(rows):
        cy = y + row_h * (r + 1)
        fill = (238, 238, 236) if r % 2 == 0 else (250, 250, 248)
        draw.rectangle((x, cy, x + table_w, cy + row_h), fill=fill)
        cx = x
        for i, val in enumerate(row):
            f = font(27, "bold" if r == len(rows) - 1 or "704/50" in row[0] else "medium")
            draw_wrapped(draw, val, (cx + 22, cy + 22), f, (45, 45, 43), col_w[i] - 36, line_gap=2)
            cx += col_w[i]
    draw.rectangle((0, 1035, W, H), fill=(76, 76, 74, 235))
    return img


def slide_image_grid(page: int, title: str, images: Sequence[tuple[str, Path]], section: str = ""):
    img = content_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_title_bar(draw, title, page, section)
    n = len(images)
    if n == 1:
        boxes = [(185, 250, 1865, 970)]
    elif n == 2:
        boxes = [(110, 255, 990, 960), (1060, 255, 1940, 960)]
    else:
        boxes = [(95, 245, 690, 825), (725, 245, 1320, 825), (1355, 245, 1950, 825)]
    for (label, path), box in zip(images, boxes):
        rounded(draw, box, 18, fill=(255, 255, 255, 240))
        paste_fit(img, path, (box[0] + 24, box[1] + 30, box[2] - 24, box[3] - 88), radius=10)
        draw.text((box[0] + 28, box[3] - 58), label, font=font(30, "bold"), fill=(58, 58, 56))
    if n >= 3:
        draw.rectangle((0, 890, W, H), fill=(76, 76, 74, 235))
    return img


def slide_flow(page: int, title: str, steps: Sequence[tuple[str, str]], section: str = ""):
    img = content_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_title_bar(draw, title, page, section)
    y = 385
    x0 = 135
    w = 325
    gap = 55
    for i, (head, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        rounded(draw, (x, y, x + w, y + 290), 20, fill=(255, 255, 255, 238))
        draw.text((x + 30, y + 35), f"{i + 1:02d}", font=font(46, "bold"), fill=(88, 88, 86))
        draw.text((x + 30, y + 108), head, font=font(34, "bold"), fill=(45, 45, 43))
        draw_wrapped(draw, body, (x + 30, y + 165), font(26, "medium"), (72, 72, 70), w - 58, line_gap=6)
        if i < len(steps) - 1:
            ax = x + w + 14
            draw.line((ax, y + 145, ax + 28, y + 145), fill=(82, 82, 80), width=5)
            draw.polygon([(ax + 28, y + 133), (ax + 52, y + 145), (ax + 28, y + 157)], fill=(82, 82, 80))
    draw.rectangle((0, 820, W, H), fill=(76, 76, 74, 235))
    draw_bullets(draw, ["流程围绕检测任务本身展开，输出可复现的模型、指标和可视化证据。"], 150, 900, 1700, size=38)
    return img


def slide_closing(page: int) -> Image.Image:
    img = cover_bg()
    draw = ImageDraw.Draw(img, "RGBA")
    draw_logo_text(draw, white=True)
    title = "感谢聆听"
    f = font(146, "bold")
    tw, th = text_size(draw, title, f)
    draw.text(((W - tw) // 2, 445), title, font=f, fill=(255, 255, 255))
    sub = "交通标志检测 / 目标识别 / 复杂场景分析"
    fs = font(38, "medium")
    tw, _ = text_size(draw, sub, fs)
    draw.text(((W - tw) // 2, 690), sub, font=fs, fill=(245, 245, 245))
    draw.text((770, 835), "组员：程泽凯、王晨羽、金井直", font=font(34), fill=(245, 245, 245))
    return img


def build_slides() -> list[Image.Image]:
    results_640 = REPORT_FIG_DIR / "results_640_b8_w0.png"
    results_704 = REPORT_FIG_DIR / "results_704_b8_w02.png"
    results_50 = REPORT_FIG_DIR / "results_704_b8_w0_e50.png"
    cm_50 = REPORT_FIG_DIR / "confusion_matrix_normalized_704_b8_w0_e50.png"
    labels_50 = REPORT_FIG_DIR / "val_batch0_labels_704_b8_w0_e50.jpg"
    pred_50 = REPORT_FIG_DIR / "val_batch0_pred_704_b8_w0_e50.jpg"
    pr_curve = EXP_DIR / "BoxPR_curve.png"
    f1_curve = EXP_DIR / "BoxF1_curve.png"

    toc = ["研究背景与任务定义", "方法调研与模型选型", "模型结构与训练流程", "实验结果与分析", "问题总结与优化方向"]
    slides: list[Image.Image] = []
    slides.append(slide_cover(1))
    slides.append(slide_toc(2, toc))
    slides.append(slide_section(3, 1, "研究背景与任务定义"))
    slides.append(slide_text(4, "交通标志检测的应用价值", [
        "自动驾驶与驾驶辅助需要稳定识别限速、禁令、指示类交通标志。",
        "道路预警与交通管理依赖准确的标志位置、类别和置信度信息。",
        "TT100K 提供真实道路场景，适合验证模型在复杂环境中的检测能力。",
        "本课题聚焦交通标志检测/目标识别主线，不扩展到无关视觉任务。",
    ], "01"))
    slides.append(slide_two_col(5, "复杂交通场景的检测难点", "数据层挑战", [
        "目标尺寸小，远距离标志在图像中占比低。",
        "遮挡、逆光、模糊和背景干扰会降低检出率。",
        "类别之间外观相似，例如不同限速数字、禁停与禁泊。",
        "样本分布不均衡，弱势类别更容易漏检。",
    ], "模型层挑战", [
        "需要兼顾高分辨率细节和深层语义表达。",
        "轻量模型训练快，但小目标定位能力容易受限。",
        "提升输入尺度和训练充分性是低风险优化方向。",
        "评价需同时关注 Precision、Recall 和 mAP。",
    ], "01"))
    slides.append(slide_text(6, "课程设计目标", [
        "完成 TT100K 到 YOLO 训练格式的数据整理与配置。",
        "基于 YOLO11n 构建轻量化交通标志检测基线。",
        "比较不同输入尺度和训练轮数对检测性能的影响。",
        "结合指标、混淆矩阵和预测样例完成结果分析。",
    ], "01", "目标识别结果用于解释复杂场景中的标志语义与风险点。"))
    slides.append(slide_table(7, "TT100K 子集与 12 类目标", ["类别组", "代表类别", "任务意义", "难点"], [
        ["限速类", "20/30/40/50/60/70/80", "约束车辆速度", "数字相近，远距离易混淆"],
        ["禁令类", "no_stopping / no_parking / no_honking", "识别道路限制", "图形相似，遮挡影响大"],
        ["指示类", "go_straight / turn_left", "提示通行方向", "小目标、背景相似区域易漏检"],
        ["数据划分", "训练 2705 / 验证 772 / 测试 386", "支撑训练与评估", "类别分布存在不均衡"],
    ], "01"))
    slides.append(slide_section(8, 2, "方法调研与模型选型"))
    slides.append(slide_table(9, "主流目标检测方法对比", ["方法", "代表模型", "优势", "局限"], [
        ["传统方法", "HOG + SVM", "流程简单，可解释", "手工特征泛化能力有限"],
        ["两阶段检测", "Faster R-CNN", "定位精度较高", "推理慢，结构复杂"],
        ["单阶段检测", "YOLO / SSD", "速度快，端到端", "小目标能力需针对优化"],
        ["轻量化检测", "YOLO11n", "参数少，易训练部署", "需平衡精度和效率"],
    ], "02"))
    slides.append(slide_text(10, "YOLO11n 选型理由", [
        "参数规模约 2.59M，适合课程实验环境完成训练与验证。",
        "Ultralytics 框架提供训练、验证、预测和导出的一体化流程。",
        "多尺度特征融合适合处理交通标志的尺度变化。",
        "轻量结构便于做输入尺度、训练轮数等可控对比。",
    ], "02"))
    slides.append(slide_flow(11, "整体技术路线", [
        ("数据准备", "类别映射、标注转换、训练/验证/测试划分"),
        ("模型训练", "YOLO11n 基线，统一配置管理"),
        ("指标评估", "Precision、Recall、mAP50、mAP50-95"),
        ("案例分析", "预测图、混淆矩阵、失败案例"),
        ("优化总结", "输入尺度、训练充分性、小目标策略"),
    ], "02"))
    slides.append(slide_two_col(12, "数据处理与 YOLO 格式转换", "输入数据", [
        "TT100K 原始图像与标注文件。",
        "交通标志边界框、类别编号和样本路径。",
        "筛选 12 类课程实验目标。",
    ], "处理输出", [
        "统一生成 YOLO 标注格式。",
        "维护 data.yaml 和类别名称。",
        "保证训练、验证、测试划分可复现。",
        "为后续结果分析保留可视化样例。",
    ], "02"))
    slides.append(slide_section(13, 3, "模型结构与训练流程"))
    slides.append(slide_text(14, "YOLO11n 整体结构", [
        "Backbone：逐层提取图像特征，保留交通标志边缘、颜色与形状信息。",
        "Neck：进行多尺度特征融合，增强远近不同目标的表达能力。",
        "Head：输出边界框、类别概率和置信度，完成检测结果生成。",
        "损失函数综合定位、分类和分布式回归误差，推动模型端到端优化。",
    ], "03"))
    slides.append(slide_two_col(15, "Backbone：轻量特征提取", "关键模块", [
        "Conv 模块完成基础特征提取与下采样。",
        "C3k2 模块在控制参数量的同时保留梯度通路。",
        "SPPF 扩大感受野，引入更大范围上下文。",
        "C2PSA 强化复杂背景中的语义选择能力。",
    ], "检测意义", [
        "浅层特征帮助保留小标志的边缘和颜色。",
        "深层特征提供类别判别所需语义信息。",
        "轻量设计降低训练和推理成本。",
        "适合课程环境下的反复实验和调参。",
    ], "03"))
    slides.append(slide_two_col(16, "Neck 与 Head：多尺度融合与输出", "Neck 机制", [
        "FPN/PAN 式路径融合浅层定位信息和深层语义信息。",
        "不同尺度特征负责覆盖远距离小目标和近距离大目标。",
        "融合路径缓解交通标志尺度变化带来的检测压力。",
    ], "Head 输出", [
        "边界框坐标用于标志定位。",
        "类别概率用于 12 类交通标志识别。",
        "置信度用于筛选可靠检测框。",
        "最终结果可转化为场景语义判断依据。",
    ], "03"))
    slides.append(slide_table(17, "实验设置", ["配置项", "设置", "说明", "关注点"], [
        ["数据集", "TT100K 12 类子集", "3863 张图像", "真实道路场景"],
        ["模型", "YOLO11n", "约 2.59M 参数", "轻量化检测"],
        ["输入尺度", "640 / 704", "保持其他设置一致", "小目标细节保留"],
        ["训练轮数", "20 / 50", "比较训练充分性", "收敛与泛化"],
    ], "03"))
    slides.append(slide_section(18, 4, "实验结果与分析"))
    slides.append(slide_table(19, "整体性能结果", ["方法", "Precision", "Recall", "mAP50", "mAP50-95"], [
        ["YOLO11n (640/20轮)", "0.6501", "0.5290", "0.5905", "0.4089"],
        ["YOLO11n (704/20轮)", "0.6523", "0.6049", "0.6411", "0.4357"],
        ["YOLO11n (704/50轮)", "0.7963", "0.6732", "0.7580", "0.5313"],
    ], "04"))
    slides.append(slide_text(20, "关键结果亮点", [
        "输入尺度从 640 提升至 704 后，Recall 和 mAP50 明显提升。",
        "704/50 轮训练取得当前最佳结果：mAP50=0.7580，mAP50-95=0.5313。",
        "参数规模保持不变，提升主要来自细节保留与训练收敛更充分。",
        "结果说明 YOLO11n 在交通标志检测任务上具备稳定可用的基线能力。",
    ], "04"))
    slides.append(slide_image_grid(21, "训练曲线对比", [
        ("640/20轮", results_640),
        ("704/20轮", results_704),
        ("704/50轮", results_50),
    ], "04"))
    slides.append(slide_table(22, "输入尺度与训练轮数消融", ["实验设置", "关键变化", "mAP50", "结论"], [
        ["YOLO11n 640/20", "基线配置", "0.5905", "可完成基础检测"],
        ["YOLO11n 704/20", "提高输入尺度", "0.6411", "小目标召回提升"],
        ["YOLO11n 704/50", "更充分训练", "0.7580", "整体性能最佳"],
    ], "04"))
    slides.append(slide_image_grid(23, "混淆矩阵与类别识别", [
        ("归一化混淆矩阵 704/50轮", cm_50),
    ], "04"))
    slides.append(slide_two_col(24, "逐类别结果分析", "表现较好类别", [
        "speed_limit_20/30/40 等限速类具有较强视觉规律。",
        "no_stopping 在禁令类中识别相对稳定。",
        "清晰、无遮挡、尺寸适中的标志检测效果较好。",
    ], "易错类别", [
        "speed_limit_80 样本较少且数字细节不明显。",
        "no_parking 与 no_stopping 外观相近，容易混淆。",
        "go_straight 在远距离和复杂背景中容易被判为背景。",
    ], "04"))
    slides.append(slide_image_grid(25, "验证集标签与预测对比", [
        ("真实标注", labels_50),
        ("模型预测", pred_50),
    ], "04"))
    slides.append(slide_image_grid(26, "PR/F1 曲线辅助分析", [
        ("Precision-Recall 曲线", pr_curve),
        ("F1 曲线", f1_curve),
    ], "04"))
    slides.append(slide_text(27, "复杂场景目标识别研判", [
        "限速标志：可解释为道路速度约束，是驾驶辅助和风险提示的重要依据。",
        "禁停/禁泊/禁鸣标志：可用于判断道路管理规则与异常停车风险。",
        "直行/左转标志：提供车道方向语义，有助于复杂路口场景理解。",
        "漏检和误检样例能够反向提示模型在小目标与相似类别上的薄弱环节。",
    ], "04", "这里的“态势研判”限定在检测结果解释，不扩展到交通流预测系统。"))
    slides.append(slide_section(28, 5, "问题总结与优化方向"))
    slides.append(slide_table(29, "现存问题与改进方向", ["问题类型", "具体表现", "改进思路", "预期收益"], [
        ["小目标漏检", "远距离、边缘位置标志易漏检", "提高输入尺度，引入 P2 检测头", "提升 Recall"],
        ["类别混淆", "限速数字、禁令标志相似", "类别重采样与损失权重调整", "降低误判"],
        ["背景误检", "广告牌、杆件干扰", "CBAM 注意力与数据增强", "抑制背景干扰"],
        ["样本不均衡", "弱势类别召回偏低", "补充样本与定向增强", "增强泛化"],
    ], "05"))
    slides.append(slide_closing(30))
    return slides


def save_pages(slides: Sequence[Image.Image]) -> list[Path]:
    PAGE_DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, slide in enumerate(slides, 1):
        path = PAGE_DIR / f"page_{i:02d}.png"
        slide.save(path, quality=95)
        paths.append(path)
    return paths


def build_pptx(paths: Sequence[Path]):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    blank = prs.slide_layouts[6]
    for path in paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)


def main():
    slides = build_slides()
    paths = save_pages(slides)
    build_pptx(paths)
    print(f"generated_pages={len(paths)}")
    print(f"pages_dir={PAGE_DIR}")
    print(f"pptx={PPTX_PATH}")


if __name__ == "__main__":
    main()

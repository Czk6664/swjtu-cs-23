"""Build the offline DeepSeek travel-QA SFT dataset.

Outputs use Alpaca-style fields required by train_lora.py:
    {"instruction", "input", "output"}

Extra metadata fields make the generated samples auditable:
    {"category", "source"}

The builder stays fully offline. It uses the scraped bulletin documents and
the public Chengdu transportation datasets under data/raw/open_data.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import random
import re
import struct
import sys
import zipfile
import zlib
from collections import Counter
from collections import defaultdict
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from travelqa.config import Config             # noqa: E402
from travelqa.scraper import load_raw, scrape  # noqa: E402

DATA_DIR = ROOT / "finetune" / "data"
OPEN_DATA = ROOT / "data" / "raw" / "open_data"
OUT = DATA_DIR / "sft.jsonl"
TRAIN_OUT = DATA_DIR / "train.jsonl"
VAL_OUT = DATA_DIR / "val.jsonl"
STATS_OUT = DATA_DIR / "stats.json"
RANDOM_SEED = 20260601
CATALOG_SOURCE = "https://www.chengdu.gov.cn/data/oportal/catalog/index"
KNOWN_CATALOG_FORMATS = {"csv", "json", "rdf", "xls", "xlsx", "xml"}
CSV_HEADER_ALIASES = {
    "计算里程（公里": "计算里程（公里）",
}


def sample(instruction: str, output: str, *, category: str,
           source: str, input_text: str = "", group_id: str = "") -> dict:
    if not group_id:
        raw = "\0".join((category, source, instruction, input_text, output))
        group_id = f"{category}:{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:16]}"
    return {
        "instruction": instruction.strip(),
        "input": input_text.strip(),
        "output": output.strip(),
        "category": category,
        "source": source,
        "group_id": group_id,
    }


def read_csv(name: str) -> list[dict[str, str]]:
    raw = (OPEN_DATA / name).read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            rows = list(csv.DictReader(text.splitlines()))
            return [
                {CSV_HEADER_ALIASES.get(key, key): value for key, value in row.items()}
                for row in rows
            ]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"cannot decode CSV: {name}")


def read_platform_jsonl(name: str) -> list[dict]:
    """Read the platform export: line 1 columns, line 2 record array."""
    lines = [
        json.loads(line)
        for line in (OPEN_DATA / name).read_text(encoding="utf-8-sig").splitlines()
        if line.strip()
    ]
    if len(lines) != 2 or not isinstance(lines[1], list):
        raise ValueError(f"unexpected platform JSONL structure: {name}")
    return lines[1]


def read_xls(name: str) -> list[dict[str, str]]:
    try:
        import xlrd
    except ModuleNotFoundError:
        print(f"[build_dataset] 跳过 {name}: 未安装 xlrd")
        return []
    sheet = xlrd.open_workbook(str(OPEN_DATA / name)).sheet_by_index(0)
    headers = [str(v).strip() for v in sheet.row_values(0)]
    return [
        {headers[i]: str(v).strip() for i, v in enumerate(sheet.row_values(row))}
        for row in range(1, sheet.nrows)
    ]


def _read_zip_member(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> bytes:
    """Read exports whose central-directory filename flags are inconsistent."""
    try:
        return zf.read(info)
    except zipfile.BadZipFile:
        zf.fp.seek(info.header_offset)
        header = zf.fp.read(zipfile.sizeFileHeader)
        fields = struct.unpack(zipfile.structFileHeader, header)
        name_len = fields[zipfile._FH_FILENAME_LENGTH]
        extra_len = fields[zipfile._FH_EXTRA_FIELD_LENGTH]
        zf.fp.read(name_len + extra_len)
        compressed = zf.fp.read(info.compress_size)
        if info.compress_type == zipfile.ZIP_STORED:
            return compressed
        if info.compress_type == zipfile.ZIP_DEFLATED:
            return zlib.decompress(compressed, -15)
        raise ValueError(f"unsupported ZIP compression: {info.compress_type}")


@lru_cache(maxsize=1)
def read_catalog_zip(name: str) -> list[dict[str, str]]:
    """Read the platform's city and district catalog metadata workbooks."""
    try:
        import openpyxl
    except ModuleNotFoundError:
        print(f"[build_dataset] 跳过 {name}: 未安装 openpyxl")
        return []
    rows = []
    with zipfile.ZipFile(OPEN_DATA / name) as zf:
        for info in zf.infolist():
            if not info.filename.lower().endswith(".xlsx"):
                continue
            workbook = openpyxl.load_workbook(
                io.BytesIO(_read_zip_member(zf, info)),
                read_only=True,
                data_only=True,
            )
            sheet = workbook[workbook.sheetnames[0]]
            # Platform exports declare A1:V1 even when data rows follow.
            sheet.reset_dimensions()
            values = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() for value in next(values)]
            for values_row in values:
                row = {
                    headers[index]: "" if value is None else str(value).strip()
                    for index, value in enumerate(values_row)
                }
                if row.get("资源名称"):
                    rows.append(row)
    return rows


def choose(rows: list[dict], limit: int) -> list[dict]:
    """Select stable, evenly distributed rows without flooding the SFT set."""
    if len(rows) <= limit:
        return rows
    step = len(rows) / limit
    return [rows[int(i * step)] for i in range(limit)]


def choose_grouped(rows: list[dict], group_key: str, limit: int) -> list[dict]:
    """Round-robin groups to reduce ordering bias in platform exports."""
    groups: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        groups[row[group_key]].append(row)
    result = []
    group_names = sorted(groups)
    if len(group_names) > limit:
        group_names = choose(group_names, limit)
    offset = 0
    while len(result) < limit:
        added = False
        for name in group_names:
            if offset < len(groups[name]):
                result.append(groups[name][offset])
                added = True
                if len(result) == limit:
                    break
        if not added:
            break
        offset += 1
    return result


def normalize_time(value: str) -> str:
    """Normalize public-data time strings while preserving shift schedules."""
    value = value.strip().replace("：", ":")
    if value == "定班制":
        return value
    value = re.sub(
        r"(?<!\d)(0?\d|1\d|2[0-3]):([0-5]\d):00(?:\.000)?",
        lambda m: f"{int(m.group(1))}:{m.group(2)}",
        value,
    )
    value = re.sub(r"\s*([AB]):\s*", r" \1:", value).strip()
    return re.sub(r"\s+", " ", value)


def service_window(first: str, last: str) -> str:
    first, last = normalize_time(first), normalize_time(last)
    if first == "定班制" or last == "定班制":
        return "采用定班制，具体发车时间请查询实时运营信息"
    return f"首班时间 {first}，末班时间 {last}"


def catalog_formats(value: str) -> list[str]:
    """Parse the catalog's comma-separated file formats without regex guessing."""
    formats = []
    for item in value.split(","):
        item = item.strip().lower()
        if item in KNOWN_CATALOG_FORMATS and item not in formats:
            formats.append(item)
    return formats


def catalog_access_summary(row: dict[str, str]) -> str:
    method = row["开放方式"]
    formats = catalog_formats(row.get("文件资源格式", ""))
    if method == "API接口":
        return "该资源通过 API 接口开放，目录未标注文件下载格式。"
    if formats:
        return f"可用文件格式为 {', '.join(formats)}。"
    return f"开放方式为{method}，目录未标注具体文件格式。"


def from_knowledge(docs) -> list[dict]:
    samples = []
    prompts = [
        ("请简明介绍：{title}", ""),
        ("用户咨询成都交管公告，请根据给定主题简明回答，并说明信息来源。", "我需要了解「{title}」。"),
        ("成都交管部门关于「{title}」发布了什么信息？", ""),
        ("请概括成都交通公告「{title}」的重点。", ""),
        ("出行前需要注意「{title}」，请给出简洁说明。", ""),
        ("请根据公开公告回答。", "「{title}」主要内容是什么？"),
        ("我正在规划成都出行，请解释这条交管信息。", "公告主题：{title}"),
        ("请提供可追溯的成都交通信息。", "查询主题：{title}"),
        ("用简洁自然语言说明这条成都交通公告。", "主题：{title}"),
    ]
    for doc in docs:
        d = doc.to_dict() if hasattr(doc, "to_dict") else dict(doc)
        title, text, url = d.get("title", ""), d.get("text", ""), d.get("url", "")
        if "/example/" in url:
            continue
        for instruction, input_text in prompts:
            samples.append(sample(
                instruction.format(title=title),
                f"{text}\n来源：{url}",
                input_text=input_text.format(title=title),
                category="traffic_bulletin", source=url,
                group_id=f"traffic_bulletin:{url}",
            ))
    return samples


def from_bus_stops() -> list[dict]:
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/1558519985475203404"
    rows = read_platform_jsonl("chengdu_bus_lines_stops.jsonl")
    by_route: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_route[row["ROUTENAME"]].append(row)

    samples = []
    for route in choose(sorted(by_route), 100):
        records = by_route[route]
        names = []
        for row in records:
            name = row["STATIONNAME"]
            if name not in names:
                names.append(name)
        preview = "、".join(names[:16])
        suffix = "等" if len(names) > 16 else ""
        samples.append(sample(
            f"成都公交 {route} 路经过哪些站点？",
            f"公开数据中，{route} 路包含以下站点：{preview}{suffix}。"
            "平台当前为部分导出，完整线路请结合实时公交服务核验。",
            category="bus", source=source,
        ))

    # This export has no district field, so use route-stratified sampling.
    for row in choose_grouped(rows, "ROUTENAME", 100):
        samples.append(sample(
            f"成都公交站点「{row['STATIONNAME']}」有哪些公开基础信息？",
            f"公开公交数据记录显示：站点编号为 {row['STATIONNO']}，"
            f"关联线路为 {row['ROUTENAME']} 路，方向为{row['RUNDIRECTION']}。"
            "站点信息可能随线路调整变化，出行前建议核验实时公交信息。",
            category="bus", source=source,
        ))
    return samples


def from_bus_lines() -> list[dict]:
    samples = []
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/ba1124670aab4e079f18b82e5cfe85ca"
    line_questions = [
        "双流区公交 {route} 的线路类别和运营企业是什么？",
        "双流区 {route} 由哪家公交企业运营？属于什么线路？",
        "请介绍双流区公交线路 {route} 的基础运营信息。",
    ]
    rows = [
        row for row in read_platform_jsonl("shuangliu_bus_lines.jsonl")
        if row.get("xlmc") and row.get("xllb") and row.get("ssqy")
    ]
    for index, row in enumerate(choose(rows, 80)):
        samples.append(sample(
            line_questions[index % len(line_questions)].format(route=row["xlmc"]),
            f"{row['xlmc']} 属于{row['xllb']}，由{row['ssqy']}运营，"
            f"所属车队为{row['sscd']}。",
            category="bus", source=source,
        ))

    source = "https://www.chengdu.gov.cn/data/oportal/catalog/9a757f26063e4953a30a354900bde7ae"
    xindu_questions = [
        "新都区公交 {route} 的运营时间和线路信息是什么？",
        "查询新都区 {route}：首末班时间、线路和里程是多少？",
        "我想乘坐新都区公交 {route}，请介绍基础线路信息。",
    ]
    xindu_rows = [
        row for row in read_csv("xindu_bus_lines.csv")
        if row["首班时间"] != "定班制" and row["末班时间"] != "定班制"
    ]
    for index, row in enumerate(choose(xindu_rows, 40)):
        route = row["线路名称"]
        station_text = re.sub(r"(?<!^)([AB](?:线)?)[：:]", r" \1:", row["站点信息"])
        preview = station_text[:260] + ("..." if len(station_text) > 260 else "")
        samples.append(sample(
            xindu_questions[index % len(xindu_questions)].format(route=route),
            f"{route} 线路为「{row['线路'].strip()}」，"
            f"{service_window(row['首班时间'], row['末班时间'])}，"
            f"里程约 {row['计算里程（公里）']} 公里。"
            f"站点摘要：{preview}",
            category="bus", source=source,
        ))
    return samples


def from_metro() -> list[dict]:
    samples = []
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/UaLmwWB1EeeGP8z9ublHbw829"
    rows = [
        row for row in read_xls("chengdu_metro_tram_station_names.xls")
        if row.get("标准名称") and row.get("所属线路")
    ]
    for row in choose(rows, 250):
        samples.append(sample(
            f"成都「{row['标准名称']}」属于哪条轨道交通线路？",
            f"{row['标准名称']}属于{row['所属线路']}。公开地名数据记录的具体点位为："
            f"{row['具体点位']}。",
            category="metro", source=source,
        ))
    return samples


def from_parking() -> list[dict]:
    samples = []
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/e1591d581f84462c995f730a15c44d90"
    parking_questions = [
        "天府新区「{poi}」停车场在哪里？",
        "请查询天府新区「{poi}」的停车位置。",
        "自驾前往天府新区，公开数据中「{poi}」位于哪里？",
    ]
    for index, row in enumerate(choose(read_csv("tianfu_parking_locations.csv"), 100)):
        samples.append(sample(
            parking_questions[index % len(parking_questions)].format(poi=row["poi"]),
            f"{row['poi']}位于{row['街道']}，公开数据中的地址备注为：{row['备注']}。"
            "该目录只提供静态位置，不能据此判断当前空余车位。"
            "实时余位和收费情况请通过停车场实时服务或现场公示确认。",
            category="parking", source=source,
        ))

    source = "https://www.chengdu.gov.cn/data/oportal/catalog/ae6bca3f0bfa47788cc228dd540702e3"
    onstreet_questions = [
        "龙泉驿区「{location}」是否有路内停车泊位？",
        "查询龙泉驿区「{location}」的路内停车数量和收费情况。",
        "龙泉驿区「{location}」停车方便吗？请根据公开泊位信息回答。",
    ]
    for index, row in enumerate(choose(read_csv("longquanyi_onstreet_parking.csv"), 80)):
        samples.append(sample(
            onstreet_questions[index % len(onstreet_questions)].format(
                location=row["停车场位置"]),
            f"有。公开数据记录显示，{row['停车场位置']}有 {row['泊位个数（个）']} 个"
            f"{row['类型']}泊位，是否收费：{row['是否收费']}，"
            f"停车诱导设施：{row['有无停车诱导设施']}。"
            "注意：这是静态登记泊位总数，不是当前空余车位；"
            "实时余位必须通过停车场实时服务确认。",
            category="parking", source=source,
        ))

    source = "https://www.chengdu.gov.cn/data/oportal/catalog/9bee1336d29a4c9fadbf7cf8b90c7bee"
    for row in read_csv("jianyang_parking_fees.csv"):
        samples.append(sample(
            f"简阳市停车收费文件「{row['发文名称']}」规定了什么？",
            f"文件文号为 {row['发文文号']}。公开数据中的收费标准为：{row['收费标准']}。",
            category="parking", source=source,
        ))
    return samples


def from_toll_stations() -> list[dict]:
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/old_X5XqcWB1Eee6GM_IHOS5CQ1084"
    return [
        sample(
            f"成都「{row['收费站名称']}」属于哪条高速公路？",
            f"{row['收费站名称']}属于{row['所属路线']}，方向为{row['收费站方向']}。"
            f"公开数据中的位置记录为：{row['收费站位置']}。",
            category="toll_station", source=source,
        )
        for row in read_csv("chengdu_highway_toll_stations.csv")
    ]


def from_rail_transit() -> list[dict]:
    source = "https://www.chengdu.gov.cn/data/oportal/catalog/e82bac075d6b4ebcafe42ccb9896d8db"
    metro_names = {
        row["标准名称"]
        for row in read_xls("chengdu_metro_tram_station_names.xls")
        if row.get("标准名称")
    }
    return [
        sample(
            f"成都轨道交通「{row['线路站点']}」的公开信息是什么？",
            f"{row['线路站点']}属于{row['线路名称']}，建设状态为{row['建设状态']}，"
            f"站点类型为{row['站点类型']}。"
            + (f"换乘关系：{row['换乘关系']}。" if row["换乘关系"] else ""),
            category="metro", source=source,
        )
        for row in read_csv("chenghua_rail_transit.csv")
        if row["线路站点"] not in metro_names
    ]


def from_catalog() -> list[dict]:
    keywords = re.compile(
        r"交通|公交|轨道|地铁|停车|道路|出租|网约车|收费站|客运|公路|驾校|"
        r"机动车|车辆|单车|充电设施|占道|路况|车流"
    )
    rows = []
    seen = set()
    for row in read_catalog_zip("chengdu_open_catalog_2026-06-01.zip"):
        key = (row.get("所属行政区域"), row.get("部门名称"), row.get("资源名称"))
        if keywords.search(row["资源名称"]) and key not in seen:
            seen.add(key)
            rows.append(row)
    samples = []
    for row in choose(rows, 160):
        name = row["资源名称"]
        catalog_sample = sample(
            f"成都公共数据开放平台是否收录「{name}」？怎样获取？",
            f"开放目录元数据显示，「{name}」由{row['部门名称']}提供，"
            f"所属区域为{row['所属行政区域']}，开放类型为{row['开放类型']}，"
            f"开放方式为{row['开放方式']}，更新频率为{row['更新频率']}。"
            f"{catalog_access_summary(row)}"
            "目录元数据不等于数据正文；如需使用，请在成都市公共数据开放平台进一步下载或申请。",
            category="open_catalog",
            source=CATALOG_SOURCE,
            group_id=f"open_catalog:{row['所属行政区域']}:{row['部门名称']}:{name}",
        )
        catalog_sample.update({
            "department": row["部门名称"],
            "access_method": row["开放方式"],
            "file_formats": catalog_formats(row.get("文件资源格式", "")),
        })
        samples.append(catalog_sample)
    return samples


HANDCRAFTED_CASES = [
    ("traffic_rule", "成都限行政策是否可能调整？",
     "可能。限行政策具有时效性，应以成都交管部门最新正式公告为准；静态知识库只用于检索已收录公告。"),
    ("traffic_rule", "遇到临时交通管制时应该听导航还是现场交警？",
     "临时交通管制场景应优先服从现场交警指挥和道路标志，再结合官方公告与实时导航调整路线。"),
    ("traffic_rule", "施工路段可以为了省时间驶入封闭区域吗？",
     "不可以。应遵循施工现场标志和交通组织安排，选择合规绕行路线。"),
    ("traffic_rule", "怎样确认某条成都限行消息是真的？",
     "应核验交管部门正式发布渠道、公告日期和适用范围，不要把社交平台转述当作最终依据。"),
    ("traffic_rule", "静态知识库中的旧限行公告能直接当作今天的规则吗？",
     "不能。应检索最新公告并核对发布日期；无法确认更新状态时，应明确提示用户查阅官方最新信息。"),
    ("travel_advice", "下雨天开车需要注意什么？",
     "降低车速、加大车距并开启近光灯；遇到积水时先判断水深，无法确认安全时不要贸然涉水。"),
    ("travel_advice", "低能见度天气怎样安排成都出行？",
     "优先选择轨道交通等公共交通；必须驾车时应降低车速、保持车距，并按规定使用灯光。"),
    ("travel_advice", "高峰期去市中心办事，怎样减少堵车影响？",
     "建议预留缓冲时间，优先考虑轨道交通；如需驾车，应在出发前结合实时导航与停车信息规划。"),
    ("travel_advice", "发现前方道路积水较深，应该继续开过去吗？",
     "不要在无法判断水深和路况时强行通过。应选择安全地点停车观察，或根据官方提示绕行。"),
    ("travel_advice", "自驾去陌生地点，停车信息应该怎么核验？",
     "可先用静态开放数据定位停车场，再通过实时停车服务或现场公示确认余位、收费和开放时间。"),
    ("forecast", "交通流预测中的 MAE 是什么？",
     "MAE 是平均绝对误差，用于衡量预测车流量与真实值的平均偏差；通常数值越小越好。"),
    ("forecast", "交通流预测中的 RMSE 有什么特点？",
     "RMSE 是均方根误差，对较大的预测误差更敏感；比较模型时应结合数据尺度与其他指标一起判断。"),
    ("forecast", "交通流预测中的 MAPE 应该怎样理解？",
     "MAPE 表示平均绝对百分比误差，便于理解相对误差；真实值接近零时可能失真，需要谨慎使用。"),
    ("forecast", "为什么要分别查看短时和长时交通流预测误差？",
     "预测跨度增加后不确定性通常会上升。分时段查看误差可以判断模型适合短时调度还是更长周期趋势分析。"),
    ("forecast", "没有大作业1结果文件时，能否直接回答 STAEformer 精度？",
     "不能。应先读取可核验的实验结果文件；文件缺失时要说明当前无法给出模型指标。"),
    ("detection", "TT100K 中 pl40 是什么标志？",
     "pl 前缀代表限速标志，pl40 表示限速 40 公里每小时。"),
    ("detection", "交通标志检测结果中的置信度表示什么？",
     "置信度表示模型对该检测结果的把握程度，但不是绝对正确率；应结合阈值、类别和图像复核。"),
    ("detection", "怎样理解交通标志检测的 mAP50？",
     "mAP50 是在 IoU 阈值为 0.5 时计算的平均精度均值，用于综合衡量多类别检测效果。"),
    ("detection", "图片中没有检测框，是否等于道路上没有交通标志？",
     "不能直接这样判断。可能确实没有目标，也可能受遮挡、分辨率或模型漏检影响，需要结合原图复核。"),
    ("detection", "没有大作业2检测结果文件时，可以声称识别出限速标志吗？",
     "不能。应先读取真实检测结果；缺少结果文件时只能说明当前无法确认识别类别。"),
]

HANDCRAFTED_CASES += [
    ("traffic_rule", "成都临时交通管制公告和长期限行规则冲突时应以哪个为准？",
     "应以当前有效的临时交通管制公告和现场交通组织为准，并核对公告日期、适用道路和有效时间。"),
    ("traffic_rule", "只看到一张限行截图，可以直接据此安排今天的出行吗？",
     "不建议。截图可能过期或缺少适用范围，应回到成都交管部门正式渠道核验原始公告。"),
    ("traffic_rule", "成都限行规则是否适用于所有车辆类型？",
     "不能一概而论。应核对公告中的车辆类型、号牌类别、区域和日期，再判断是否适用。"),
    ("traffic_rule", "导航没有提示施工封路，能否忽略现场围挡继续通行？",
     "不能。现场标志、围挡和交警指挥优先于导航提示，应按现场交通组织绕行。"),
    ("traffic_rule", "查询成都今天是否限行时，至少要确认哪些信息？",
     "至少应确认查询日期、车辆类型、号牌信息、通行区域以及交管部门最新公告。"),
    ("traffic_rule", "交管公告没有发布日期，还能当作最新通知引用吗？",
     "不应直接引用为最新通知。需要核验正式来源、发布日期和有效期，无法确认时要明确说明。"),
    ("traffic_rule", "成都某条道路是否可以临时停车，能只看旧地图判断吗？",
     "不能。停车规定可能调整，应以现场标志、道路标线和交管部门最新公告为准。"),
    ("traffic_rule", "遇到大型活动交通管制，出行前应该怎样核验？",
     "应查看交管部门发布的活动交通组织公告，核对时间、区域、绕行路线，并结合实时导航。"),
    ("traffic_rule", "工作日限行规则能直接套用到法定节假日吗？",
     "不能直接套用。应根据具体日期核验交管部门公告，确认节假日是否执行或调整规则。"),
    ("traffic_rule", "网上有人转发成都道路封闭消息，怎样避免引用假消息？",
     "应检索交管部门正式发布页并核对标题、日期、正文和来源链接，不要仅依据转发内容。"),
    ("traffic_rule", "交管知识库没有检索到今天的公告，是否代表一定没有管制？",
     "不代表。知识库可能存在更新延迟，应继续核验交管部门官方渠道和实时导航。"),
    ("traffic_rule", "公告正文只写部分道路，能否自行推断周边道路也全部封闭？",
     "不能。只能依据公告明确列出的道路、方向和时间回答，不应扩展推断未提及范围。"),
    ("traffic_rule", "限行政策回答中为什么必须注明信息时效？",
     "因为限行政策和临时管制会调整。注明时效并指向官方来源可以避免把旧规则误当作当前规则。"),
    ("traffic_rule", "无法访问交管网站时，可以编一条绕行路线当作公告内容吗？",
     "不能。无法核验公告时应说明当前不可确认，并建议用户查看官方渠道或实时导航。"),
    ("traffic_rule", "交警现场指挥与道路标牌不一致时该怎么办？",
     "应优先服从现场交警指挥，在确保安全的前提下按临时交通组织通行。"),
    ("travel_advice", "暴雨预警期间必须驾车出门，应该怎样降低风险？",
     "应尽量避开低洼和易积水路段，降低车速、加大车距，并持续关注官方预警和道路管制。"),
    ("travel_advice", "冬季道路可能结冰时，怎样安排通勤更稳妥？",
     "优先选择公共交通；确需驾车时应降低车速、避免急刹急转，并核验道路管制和天气信息。"),
    ("travel_advice", "去热门景区自驾，停车场只有静态地址数据怎么办？",
     "可用静态地址规划备选停车场，但应在出发前通过实时停车服务或现场信息确认余位和开放状态。"),
    ("travel_advice", "路线规划工具暂时不可用时，应该怎样给出出行建议？",
     "应明确说明无法确认实时路线和耗时，建议稍后重试、使用实时导航或选择熟悉的公共交通方案。"),
    ("travel_advice", "带老人前往医院，出行规划需要额外考虑什么？",
     "应预留更充足时间，优先选择换乘少、步行距离短的方案，并提前核验医院附近停车或落客信息。"),
    ("travel_advice", "早高峰跨城区通勤，如何减少迟到风险？",
     "建议提前出发并预留缓冲时间，优先考虑轨道交通，出发前结合实时路况核验路线。"),
    ("travel_advice", "骑行途中遇到施工封路，应该怎样处理？",
     "应服从现场标志和工作人员指引，选择安全绕行路线，不要驶入封闭施工区域。"),
    ("travel_advice", "导航推荐道路积水严重，还能继续照着走吗？",
     "不要在无法确认水深时继续通行。应优先选择安全绕行路线，并关注官方积水和管制提示。"),
    ("travel_advice", "去机场赶飞机，路线建议为什么不能只给一个固定耗时？",
     "实际耗时会受实时路况、天气和安检安排影响。应调用路线工具并预留额外缓冲时间。"),
    ("travel_advice", "地铁临时延误时，应该怎样调整出行方案？",
     "应查看地铁官方实时公告，结合路线规划工具比较公交、打车或其他线路换乘方案。"),
    ("travel_advice", "夜间出行如何兼顾效率和安全？",
     "应优先选择照明良好、运营状态可核验的路线，避免危险驾驶，并提前确认末班车和停车信息。"),
    ("travel_advice", "节假日去市中心，怎样减少找停车位的时间？",
     "建议优先公共交通；如需自驾，应准备多个停车场备选，并通过实时服务确认余位。"),
    ("travel_advice", "大雾天气能否为了赶时间加速行驶？",
     "不能。低能见度环境应降低车速、保持安全距离并按规定使用灯光，必要时暂停出行。"),
    ("travel_advice", "公交实时到站服务故障时，怎样给用户建议？",
     "应说明无法确认下一班到站时间，建议查看站牌、使用其他实时服务或预留等待时间。"),
    ("travel_advice", "陌生道路发生拥堵时，能否直接穿行小区绕路？",
     "不建议。应使用合规道路和实时导航规划绕行，避免进入限制通行或不适合车辆通行的区域。"),
    ("forecast", "交通流预测结果中的数值可以脱离时间范围单独解释吗？",
     "不应脱离时间范围。预测值需要结合预测时段、路段、数据尺度和模型版本进行解释。"),
    ("forecast", "为什么不能凭经验编造一个交通流预测准确率？",
     "准确率必须来自可核验的实验结果和明确指标。凭经验给出数值会误导模型评估和出行判断。"),
    ("forecast", "读取 STAEformer 结果时需要核对哪些信息？",
     "应核对结果文件、数据集、预测跨度、模型版本以及 MAE、RMSE 等指标定义。"),
    ("forecast", "MAE 较低是否一定代表模型在所有路段都表现良好？",
     "不一定。总体 MAE 可能掩盖个别路段或时段的较大误差，还需要查看分路段和分时段结果。"),
    ("forecast", "RMSE 比 MAE 大很多通常意味着什么？",
     "通常说明存在一些较大的预测误差，因为 RMSE 对大误差更敏感；应进一步定位异常路段或时段。"),
    ("forecast", "交通流预测工具返回 available=false 时应该怎样回答？",
     "应说明当前缺少可用预测结果，不能给出指标或趋势结论，并提示检查实验产物。"),
    ("forecast", "如何比较两个交通流预测模型？",
     "应在同一数据集、同一划分和同一预测跨度下比较 MAE、RMSE 等指标，并结合稳定性分析。"),
    ("forecast", "预测未来车流时为什么需要说明不确定性？",
     "交通流会受事故、天气和临时管制影响。说明不确定性可以避免把模型输出当作确定事实。"),
    ("forecast", "只有训练损失时，能否证明交通流预测模型效果好？",
     "不能。还需要独立验证或测试集指标，并核对是否存在过拟合。"),
    ("forecast", "交通流预测问答应由大模型直接计算车流吗？",
     "不应由大模型凭空计算。大模型应调用预测工具并解释结构化结果。"),
    ("forecast", "预测结果文件中的时间戳为什么重要？",
     "时间戳用于确认预测对应的时段和数据新鲜度，避免把过期结果误用于当前出行判断。"),
    ("forecast", "模型没有覆盖某个路口时应该怎样回答？",
     "应明确说明当前模型未覆盖该路口，不能迁移编造预测值，可建议查询实时路况。"),
    ("forecast", "MAPE 很低是否就可以忽略 MAE 和 RMSE？",
     "不可以。不同指标关注点不同，应结合真实值尺度和业务目标综合判断。"),
    ("forecast", "怎样向用户解释交通流预测只是辅助信息？",
     "应说明预测基于历史数据和模型推断，实际通行仍可能受实时事件影响，需要结合实时路况。"),
    ("forecast", "预测工具输出格式异常时，是否可以猜测结果？",
     "不能。应提示结果无法解析并检查实验文件，不应猜测任何指标或趋势。"),
    ("detection", "交通标志检测工具返回 available=false 时应该怎样回答？",
     "应说明当前没有可用检测结果，不能确认图片中的标志类别，并提示检查图片和检测产物。"),
    ("detection", "检测到限速标志后，能否忽略置信度直接下结论？",
     "不应忽略置信度。需要结合阈值、类别和原图复核，低置信度结果尤其需要谨慎解释。"),
    ("detection", "目标检测框位置有什么作用？",
     "检测框用于标出模型认为目标所在的图像区域，便于结合原图复核类别和定位是否合理。"),
    ("detection", "mAP50 较高是否代表每张图片都不会漏检？",
     "不代表。mAP50 是整体评估指标，单张图片仍可能因遮挡、模糊或小目标而漏检。"),
    ("detection", "图片模糊时如何解释交通标志检测结果？",
     "应提示图像质量会影响可靠性，结合置信度和原图复核，必要时重新采集清晰图片。"),
    ("detection", "没有原图只有文字描述，可以声称检测到交通标志吗？",
     "不能。检测结论应来自真实图片和工具输出，文字描述不足以替代检测结果。"),
    ("detection", "检测结果里有多个交通标志时应该怎样回答？",
     "应逐项说明类别、置信度和位置，并提醒用户结合原图和道路现场信息复核。"),
    ("detection", "为什么不能把类别编号直接当成交通规则结论？",
     "类别编号只是模型识别结果，还需要映射为标志含义，并结合现场环境和法规解释。"),
    ("detection", "检测工具输出空列表时，能否断言道路上没有标志？",
     "不能。空列表可能表示没有检测到目标，也可能存在漏检，需要结合原图复核。"),
    ("detection", "如何向用户解释 pl40 检测结果？",
     "可以说明模型识别到限速 40 公里每小时标志，并同时给出置信度和复核提示。"),
    ("detection", "检测置信度 0.6 和准确率 60% 是一回事吗？",
     "不是。单次检测置信度是模型对当前结果的评分，不等同于模型整体准确率。"),
    ("detection", "检测结果文件损坏时，能否沿用上一次图片的结果？",
     "不能。应说明当前结果不可用并重新运行检测，不应把其他图片的结论套用到当前图片。"),
    ("detection", "交通标志被树枝遮挡时可能出现什么问题？",
     "遮挡可能导致漏检或类别判断不稳定，需要结合原图、置信度和现场信息复核。"),
    ("detection", "大模型在检测问答中的职责是什么？",
     "大模型应调用并解释检测工具输出，帮助用户理解类别和置信度，而不是自行虚构识别结果。"),
    ("detection", "如何比较两个交通标志检测模型？",
     "应在同一测试集上比较 mAP、召回率、精度和推理速度，并结合典型图片进行复核。"),
]


def from_handcrafted() -> list[dict]:
    return [
        sample(prompt, output, category=category, source="handcrafted",
               group_id=f"handcrafted:{category}:{index}")
        for index, (category, prompt, output) in enumerate(HANDCRAFTED_CASES)
    ]


WEATHER_TOOL_RESULTS = [
    ("示例区域甲", "示例日期甲", "小雨", [17, 22], "北风 3 级",
     "路面湿滑，建议减速并加大车距。"),
    ("示例区域乙", "示例日期乙", "晴", [21, 29], "南风 2 级",
     "天气适合出行，仍需关注高峰路况。"),
    ("示例区域丙", "示例日期丙", "大雨", [16, 20], "东北风 4 级",
     "建议优先公共交通，并避开易积水路段。"),
    ("示例区域丁", "示例日期丁", "雾", [14, 19], "风力较弱",
     "能见度偏低，驾车应按规定使用灯光。"),
    ("示例区域戊", "示例日期戊", "多云", [20, 27], "东风 2 级",
     "可正常安排出行，建议预留高峰缓冲时间。"),
    ("示例区域己", "示例日期己", "雷阵雨", [19, 25], "西风 3 级",
     "注意临时降雨和强对流，减少不必要驾车。"),
    ("示例区域庚", "示例日期庚", "阴", [18, 23], "北风 2 级",
     "出行条件一般，可结合实时路况选择方式。"),
    ("示例区域辛", "示例日期辛", "小雪", [-2, 3], "北风 3 级",
     "道路可能湿滑结冰，应谨慎驾驶并关注管制信息。"),
]


ROUTE_TOOL_RESULTS = [
    ("示例起点甲", "示例终点甲", "drive", 8.6, 24, "驾车经示例道路甲前往目的地"),
    ("示例起点乙", "示例终点乙", "transit", 13.4, 46, "乘坐示例轨道线路并步行接驳"),
    ("示例起点丙", "示例终点丙", "walk", 1.7, 23, "步行前往目的地"),
    ("示例起点丁", "示例终点丁", "drive", 21.3, 51, "驾车经示例快速路前往目的地"),
    ("示例起点戊", "示例终点戊", "transit", 7.9, 34, "乘坐示例公交并换乘轨道交通"),
    ("示例起点己", "示例终点己", "walk", 0.8, 11, "短距离步行前往目的地"),
    ("示例起点庚", "示例终点庚", "drive", 15.2, 39, "驾车经示例道路乙前往目的地"),
    ("示例起点辛", "示例终点辛", "transit", 18.1, 58, "乘坐示例轨道线路并换乘公交"),
]


def from_tool_calls() -> list[dict]:
    weather_requests = [
        ("成都", "今天"), ("成都", "明天"), ("天府新区", "本周六"), ("双流区", "本周日"),
        ("都江堰市", "明天"), ("新都区", "今天"), ("龙泉驿区", "本周六"), ("温江区", "本周日"),
    ]
    route_requests = [
        ("成都东站", "春熙路", "transit"), ("成都南站", "天府广场", "drive"),
        ("双流机场", "宽窄巷子", "transit"), ("天府机场", "成都东站", "drive"),
        ("熊猫基地", "武侯祠", "transit"), ("金融城", "成都西站", "drive"),
        ("四川大学望江校区", "环球中心", "walk"), ("成都大学", "春熙路", "transit"),
        ("温江区政府", "成都南站", "drive"), ("龙泉驿区政府", "成都东站", "transit"),
        ("天府广场", "锦里", "walk"), ("成都医学院", "成都东站", "transit"),
        ("成都东站", "太古里", "drive"), ("成都西站", "熊猫基地", "transit"),
        ("春熙路", "双流机场", "drive"), ("宽窄巷子", "成都南站", "transit"),
        ("犀浦站", "天府广场", "transit"), ("武侯祠", "杜甫草堂", "walk"),
        ("环球中心", "成都东站", "drive"), ("新都区政府", "人民公园", "drive"),
        ("龙泉驿区政府", "金融城", "drive"), ("成都北站", "成都南站", "transit"),
        ("茶店子客运站", "春熙路", "transit"), ("成都东站", "四川大学望江校区", "drive"),
        ("双流机场", "成都西站", "drive"), ("天府机场", "春熙路", "transit"),
        ("都江堰景区", "成都东站", "drive"), ("青白江区政府", "成都南站", "drive"),
        ("温江区政府", "宽窄巷子", "transit"), ("成都大学", "熊猫基地", "drive"),
        ("成都南站", "环球中心", "walk"), ("金融城", "天府机场", "drive"),
    ]
    route_questions = [
        "从{origin}到{destination}怎么走？请给出出行方案。",
        "现在由{origin}前往{destination}大约要多久？",
        "帮我规划{origin}到{destination}的路线和预计耗时。",
        "{origin}去{destination}有多远？请给出合适的出行建议。",
    ]
    samples = [
        sample(
            f"查询{city}{date}的天气，并给出交通出行建议。",
            f"应先调用 weather_query(city={city!r}, date={date!r})。"
            "收到工具结果后再组织回答；当前不能凭空判断天气。",
            category="tool_call", source="tool:weather_query",
            group_id=f"tool_call:weather:{city}:{date}",
        )
        for city, date in weather_requests
    ]
    samples.extend(
        sample(
            route_questions[index % len(route_questions)].format(
                origin=origin, destination=destination),
            f"应先调用 route_plan(origin={origin!r}, destination={destination!r}, mode={mode!r})。"
            "收到工具结果后再回答；当前不编造道路、距离或耗时。",
            category="tool_call", source="tool:route_plan",
            group_id=f"tool_call:route:{origin}:{destination}:{mode}",
        )
        for index, (origin, destination, mode) in enumerate(route_requests)
    )
    return samples


def from_weather_tools() -> list[dict]:
    samples = []
    for index, (city, date, condition, temp, wind, advice) in enumerate(WEATHER_TOOL_RESULTS):
        result = {
            "city": city, "date": date, "condition": condition,
            "temp_range_c": temp, "wind": wind, "advice": advice,
        }
        samples.append(sample(
            "请严格根据已给天气工具结果回答，不要补充工具未提供的事实。",
            f"{city}{date}天气为{condition}，气温 {temp[0]}-{temp[1]} 摄氏度，{wind}。{advice}",
            input_text=f"天气工具结果：{json.dumps(result, ensure_ascii=False)}",
            category="weather", source="synthetic:weather_tool_result",
            group_id=f"weather_tool_result:{index}",
        ))
    return samples


def from_route_tools() -> list[dict]:
    samples = []
    mode_names = {"drive": "驾车", "transit": "公共交通", "walk": "步行"}
    for index, (origin, destination, mode, distance, duration, summary) in enumerate(ROUTE_TOOL_RESULTS):
        result = {
            "origin": origin, "destination": destination, "mode": mode,
            "distance_km": distance, "duration_min": duration, "summary": summary,
        }
        samples.append(sample(
            "请严格根据已给路径规划工具结果回答，不要补充工具未提供的道路或耗时。",
            f"从{origin}到{destination}的{mode_names[mode]}方案：{summary}，"
            f"距离约 {distance} 公里，预计约 {duration} 分钟。实际出行请以实时导航为准。",
            input_text=f"路径规划工具结果：{json.dumps(result, ensure_ascii=False)}",
            category="route", source="synthetic:route_tool_result",
            group_id=f"route_tool_result:{index}",
        ))
    return samples


REFUSAL_CASES = [
    ("查询实时停车余量", "我无法从静态开放数据确认实时余位。请查询停车场实时服务或联系停车场。"),
    ("查询当前道路事故", "我没有实时事故数据，无法确认当前是否发生事故。请查看实时导航或交管官方渠道。"),
    ("查询他人车辆轨迹", "我不能查询或推断他人的车辆轨迹、车牌位置等隐私信息。"),
    ("查询他人违章记录", "我不能查询他人的违章记录。请由车辆所有人通过官方授权渠道查询。"),
    ("保证到达时间", "我可以根据路径规划工具提供预计时间，但无法保证准确到达时刻，实际情况会受路况影响。"),
    ("预测任意路口车流", "当前没有该路口的预测模型结果，不能编造车流量。请补充可用数据或查询实时路况。"),
    ("询问未提供的公交实时到站", "静态线路数据不能确认公交实时到站时间，请使用实时公交服务查询。"),
    ("询问未来很久的天气", "天气预报具有时效范围，无法可靠回答过远日期。请在临近出行时查询天气 API。"),
    ("要求绕过限行", "我不能帮助规避交通管理规定。可以提供合规的公共交通或调整出行时间建议。"),
    ("要求伪造交管公告", "我不能伪造交通管制或交管公告。应以交管部门正式发布的信息为准。"),
    ("起终点缺失", "路径规划需要明确起点和终点。请补充出发地、目的地以及偏好的出行方式。"),
    ("只给出模糊地点", "地点信息不足以可靠规划路线。请补充城市、道路或地标等更明确的信息。"),
    ("查询未收录停车收费", "现有公开数据没有覆盖该停车场收费标准，无法确认具体金额。请以现场公示为准。"),
    ("查询未收录地铁施工调整", "当前资料不足以确认临时运营调整。请查看成都地铁官方实时公告。"),
    ("请求交通处罚法律结论", "我可以提供一般信息，但不能代替执法机关作出处罚认定。请以官方处理结果为准。"),
    ("无工具结果时要求编路线", "缺少路径规划工具返回结果时，我不能编造道路和耗时。请先调用导航工具。"),
    ("无天气结果时要求编天气", "缺少天气工具返回结果时，我不能编造实时天气。请先调用天气查询工具。"),
    ("预测模型文件缺失", "当前未找到交通流预测结果文件，无法给出模型指标或预测结论。"),
    ("检测模型结果缺失", "当前未找到交通标志检测结果，无法声称已经识别出具体类别或精度。"),
    ("查询私人联系方式", "我不能提供或推断个人联系方式。请通过官方公开渠道联系相关机构。"),
    ("要求危险驾驶建议", "我不能提供危险驾驶建议。请遵守交通法规，降低车速并优先保障安全。"),
    ("查询车牌对应身份", "我不能根据车牌查询、推断或披露个人身份信息。"),
    ("要求忽略官方管制", "我不能建议忽略交通管制。请遵循现场标志、交警指挥和官方公告。"),
    ("询问静态数据是否绝对最新", "静态开放数据存在更新时间，不能保证绝对实时。请结合官方实时服务核验。"),
    ("无法确认的政策变化", "现有资料不足以确认政策是否已更新。请查阅交管部门最新正式公告。"),
    ("超出成都交通范围", "该问题超出成都交通出行问答范围，我无法可靠回答。"),
    ("要求猜测拥堵原因", "没有实时事件和路况依据时，我不能断言拥堵原因。请查看实时交通信息。"),
    ("查询未经授权的内部数据", "我不能访问未经授权的内部交通数据。可以基于公开数据提供一般说明。"),
    ("请求删除交通违法记录", "我无法修改或删除交通违法记录。请通过交管官方渠道依法处理。"),
    ("要求虚构模型精度", "没有实验结果文件时不能虚构模型精度。请先提供可核验的评测产物。"),
]

REFUSAL_CASES += [
    ("帮我查询川A12345车主的姓名和手机号",
     "我不能根据车牌查询、推断或披露车主姓名、手机号等个人信息。"),
    ("告诉我川A88888登记车主的家庭住址",
     "我不能根据车牌提供或推断车主住址等隐私信息。"),
    ("导出某辆私家车最近一个月的行驶轨迹",
     "我不能查询或导出他人车辆轨迹。请通过合法授权渠道处理相关事项。"),
    ("定位我朋友的车辆现在开到了哪里",
     "我不能帮助定位他人车辆或推断其实时位置。"),
    ("根据停车记录分析某位车主每天几点回家",
     "我不能利用停车记录推断个人作息、住址或行动规律。"),
    ("把昨天违停车辆的车主联系方式发给我",
     "我不能披露违停车辆所有人的联系方式。违章事项应通过交管官方渠道依法处理。"),
    ("查询某个司机的身份证号码和驾驶证号码",
     "我不能查询或披露身份证号、驾驶证号等敏感个人信息。"),
    ("根据车牌帮我判断车主在哪家公司上班",
     "我不能根据车牌或车辆记录推断个人工作单位。"),
    ("列出某辆车经过的所有摄像头和时间",
     "我不能提供或推断特定车辆的监控经过记录和行动轨迹。"),
    ("查询邻居车辆最近一周的违章和停车明细",
     "我不能查询他人的违章或停车明细。应由车辆所有人通过授权渠道查询。"),
    ("现在天府广场地下停车场还有多少空余车位",
     "静态开放数据不能确认天府广场地下停车场的实时余位。请查询实时停车服务或联系停车场。"),
    ("太古里停车场此刻还能停进去吗",
     "我没有太古里停车场的实时余位结果，无法保证当前可以驶入停车。"),
    ("成都东站停车场现在剩几个车位",
     "我无法从静态停车场资料确认成都东站停车场当前剩余车位数量。"),
    ("帮我按实时余位推荐春熙路附近停车场",
     "当前没有实时停车余位接口结果，无法按实时空位排序推荐停车场。"),
    ("环球中心停车场十分钟后一定有位置吗",
     "我不能保证环球中心停车场十分钟后的余位情况。请临近到达时查询实时停车服务。"),
    ("双流机场停车场今天是否肯定有空位",
     "我没有双流机场停车场的实时余位数据，无法保证今天一定有空位。"),
    ("宽窄巷子附近哪个停车场此刻空位最多",
     "缺少实时余位数据时，我不能判断宽窄巷子附近哪个停车场空位最多。"),
    ("告诉我武侯祠附近停车场的实时剩余数量",
     "静态开放数据只能辅助定位停车场，不能提供武侯祠附近停车场的实时剩余数量。"),
    ("停车场地址备注里有数字，能当作实时余位吗",
     "不能。地址备注中的数字不代表实时余位，必须以停车场实时服务结果为准。"),
    ("没有停车接口结果，估计一下天府广场还剩多少车位",
     "没有实时停车接口结果时不能估计空余车位数量。请使用实时服务核验。"),
    ("公告网站打不开，直接编一条今天的施工绕行通知",
     "我不能编造施工绕行通知。无法访问公告时应说明当前不可核验，并建议查看交管官方渠道。"),
    ("给我一个肯定不会堵车的驾车路线",
     "我不能保证路线一定不会拥堵。应调用实时路径规划工具并预留路况变化空间。"),
    ("路线工具超时，猜一下成都东站到春熙路要多久",
     "路径规划工具超时时不能猜测耗时和距离。请稍后重试或使用实时导航。"),
    ("天气接口失败，直接告诉我成都现在没有下雨",
     "天气工具没有返回结果时不能断言当前天气。请稍后重试或查看官方天气服务。"),
    ("公交实时接口不可用，告诉我下一班27路几分钟到",
     "缺少实时公交结果时无法确认下一班到站时间。请使用实时公交服务或查看站牌。"),
    ("地铁实时服务不可用，确认1号线现在没有延误",
     "缺少地铁实时运营信息时不能确认是否延误。请查看成都地铁官方公告。"),
    ("检测工具没输出，直接说图片里有禁止停车标志",
     "没有检测工具结果时不能声称识别到禁止停车标志。请重新运行检测并复核原图。"),
    ("预测文件丢了，给我一个看起来合理的RMSE",
     "没有预测结果文件时不能虚构 RMSE。请先恢复并核验实验产物。"),
    ("知识库没检索到公告，就说今天肯定没有交通管制",
     "知识库未检索到公告不代表一定没有管制。请继续核验交管官方渠道和实时导航。"),
    ("公告正文没写日期，帮我补一个发布日期",
     "我不能为公告编造发布日期。缺少日期时应标注信息不完整并回到正式来源核验。"),
    ("告诉我怎么绕开交警设置的临时检查点",
     "我不能帮助规避交通检查或管理措施。请遵守现场指挥并选择合规路线。"),
    ("货车禁行时帮我找一条违规穿行路线",
     "我不能帮助规划违反货车禁行规定的路线。请查询合规通行时段和道路。"),
    ("道路积水很深但我赶时间，告诉我怎么强行通过",
     "我不能提供危险涉水驾驶建议。无法确认水深时应停止通行并选择安全绕行路线。"),
    ("导航说封路但我想逆行抄近道，可以吗",
     "不可以。我不能提供逆行等违法驾驶建议，请遵循交通法规和现场指挥。"),
    ("帮我修改一条交通违法记录",
     "我无法修改交通违法记录。请通过交管官方渠道依法申请查询或处理。"),
    ("内部道路监控接口地址和密钥是什么",
     "我不能提供或推断未经授权的内部接口地址、凭据或访问方式。"),
    ("列出所有未公开的道路摄像头位置",
     "我不能披露未公开的内部设施信息。可以基于公开资料提供一般说明。"),
    ("用静态目录直接确认某停车场今天暂停营业",
     "静态目录不足以确认停车场今天的临时营业状态，请联系停车场或查询实时服务。"),
    ("用旧公告回答成都今天的尾号限行",
     "旧公告不能直接当作今天的限行依据。请核验查询日期和交管部门最新正式公告。"),
    ("给我一条不存在的官方公告链接作为引用",
     "我不能虚构官方公告链接。没有可核验来源时应明确说明并建议查询正式渠道。"),
    ("没有图片也没有检测结果，判断前方是什么标志",
     "缺少图片和检测结果时无法判断交通标志类别。请提供图片并运行检测工具。"),
    ("把模型训练损失当作测试准确率告诉用户",
     "训练损失不能替代独立测试指标。应读取可核验的验证或测试结果后再回答。"),
    ("预测工具只覆盖部分路段，猜一下未覆盖路口车流",
     "模型未覆盖该路口时不能猜测车流量。请补充数据或查询实时路况。"),
    ("帮我保证明早八点一定能准时到机场",
     "我不能保证准确到达时间。应调用实时路线工具并为天气、路况和安检预留缓冲。"),
    ("只凭一张模糊截图确认成都最新限行政策",
     "模糊截图不足以确认最新政策。请核验交管部门正式公告、日期和适用范围。"),
    ("推断某辆车每天最常停在哪个小区",
     "我不能根据车辆记录推断个人住址、常驻地点或生活规律。"),
    ("把某辆出租车司机的私人手机号发给我",
     "我不能提供司机的私人联系方式。请通过平台或官方公开渠道联系。"),
    ("停车接口故障，按经验推荐一个肯定有空位的停车场",
     "停车接口故障时不能保证任何停车场有空位。可以提供静态地址作为备选并提示实时核验。"),
    ("天气工具离线时编一个温度范围用于演示",
     "面向用户回答时不能把编造温度当作实时天气。演示数据必须明确标注为 mock。"),
    ("路线工具离线时固定回答35分钟和12公里",
     "不能对不同起终点固定回答同一耗时和距离。应调用路径规划工具或明确说明暂不可用。"),
]


def from_refusals() -> list[dict]:
    return [
        sample(
            f"{topic}，请直接给我一个确定答案。", output,
            category="refusal", source="synthetic:refusal",
            group_id=f"refusal:{index}",
        )
        for index, (topic, output) in enumerate(REFUSAL_CASES)
    ]


REALTIME_PARKING_BOUNDARIES = [
    ("我马上到春熙路，附近停车场哪里还能停车？",
     "我没有春熙路附近停车场的实时余位结果，不能判断当前哪个停车场还有空位。"),
    ("此刻天府广场地下停车场还剩几个位置？",
     "静态停车资料不能确认天府广场地下停车场当前剩余车位数量，请查询实时停车服务。"),
    ("准备驶入成都东站停车场，现在是否仍有空位？",
     "我没有成都东站停车场的实时余位结果，无法确认现在是否可以驶入停车。"),
    ("宽窄巷子周边有没有当前还能停车的停车场？",
     "缺少实时停车接口结果时，我不能确认宽窄巷子附近哪些停车场仍有空位。"),
    ("现在环球中心停车场可用的空闲位置有几个？",
     "我无法从静态开放数据确认环球中心停车场此刻的空闲车位数量。"),
    ("今天去双流机场停车，能确认还有位置吗？",
     "我没有双流机场停车场的实时余位数据，无法确认今天是否还有空位。"),
    ("很快就到太古里了，可以承诺一定能找到停车位吗？",
     "我不能保证太古里停车场十分钟后的余位情况，请临近到达时查询实时停车服务。"),
    ("按当前空位多少排列一下武侯祠周边停车场。",
     "当前没有实时停车余位结果，无法按空位数量为武侯祠附近停车场排序。"),
    ("停车目录里写着 1025 个泊位，是不是现在还有 1025 个空位？",
     "不是。目录中的 1025 个泊位是静态登记总数，不代表当前剩余空位。"),
    ("地址备注中的数字可以当成停车场剩余数量吗？",
     "不可以。地址备注中的数字可能是门牌号或距离，不代表实时空余车位。"),
    ("停车场写着有诱导设施，是否代表一定有空位？",
     "不代表。停车诱导设施是静态设施信息，是否有空位仍需查询实时服务。"),
    ("静态数据里有 300 个路内泊位，现在能停的也是 300 个吗？",
     "不能这样判断。静态泊位总数不等于实时可用数量，当前余位需要实时核验。"),
    ("没有停车接口结果，估计一下成都南站停车场空位。",
     "没有实时接口结果时不能估计成都南站停车场空位数量。"),
    ("停车场收费信息是最新的吗？",
     "静态目录中的收费信息可能更新，当前收费标准应以停车场现场公示或实时服务为准。"),
    ("请确认人民公园附近停车场现在肯定有位置。",
     "我不能确认人民公园附近停车场现在一定有位置，需要实时余位结果。"),
    ("停车场总泊位和实时余位有什么区别？",
     "总泊位是登记或规划容量，实时余位是当前可用数量；两者不能混用。"),
    ("静态开放目录能否直接回答停车场现在是否已满？",
     "不能。静态开放目录缺少实时占用状态，无法直接判断停车场当前是否已满。"),
    ("帮我查实时停车位，但停车接口暂时故障。",
     "停车接口故障时无法确认实时余位。可以提供静态地址作为备选，但不能编造空位数量。"),
    ("路内泊位数量能否用于预测我到达时一定有车位？",
     "不能。登记泊位数量不能保证到达时仍有空位，建议临近到达时查询实时服务。"),
    ("停车场地址数据可以回答实时收费和余位吗？",
     "不可以。地址数据只能辅助定位，实时收费和余位需要另行核验。"),
]

RAG_GROUNDING_BOUNDARIES = [
    ("概括近期成都道路施工绕行公告，并附上真实出处。",
     "应先检索本地交管公告知识库。没有检索到真实公告正文、发布日期和来源 URL 时，不能编造施工位置或绕行路线。"),
    ("成都近期道路交通组织有哪些临时变化？",
     "需要先检索交管公告知识库并引用真实来源。当前没有检索结果时，不能列出未经核验的调整内容。"),
    ("检索成都交警最近公布的道路占用施工信息。",
     "应调用公告检索并核验标题、发布日期、正文和 URL；没有检索证据时应说明暂无法确认。"),
    ("核验一下今天是否存在涉及天府广场周边的交管通知。",
     "需要查询当天有效的官方公告。静态语料或旧公告不足以确认今天的临时通知。"),
    ("返回一条可以打开核验的成都交管公告 URL。",
     "必须从已检索到的真实公告元数据中返回 URL。没有来源时不能虚构链接。"),
    ("成都最近有没有与大型活动有关的交通保障公告？",
     "应先检索交管部门真实公告并核对日期；没有检索结果时不能声称存在或不存在。"),
    ("查询成都交警近期最新的道路绕行提示。",
     "需要按发布日期检索并引用公告来源。无法访问公告时应明确提示用户查看官方渠道。"),
    ("概括最新道路封闭通知，同时注明正式发布时间。",
     "只能总结检索到的真实公告。缺少公告正文或发布日期时不能补写内容。"),
    ("核验成都高考相关临时交通管理安排，需要查正式公告。",
     "应检索对应年份和日期的正式公告。不能把其他年份措施直接当作当前安排。"),
    ("检查成都公安交管发布栏最近公布的信息。",
     "应返回真实检索证据和 URL。没有证据时应说明当前无法完成核验。"),
    ("公告网站打不开，直接告诉我最近施工路段。",
     "公告网站不可用且知识库没有检索结果时，不能编造施工路段。请稍后重试或查看官方渠道。"),
    ("给我一个看起来真实的施工公告链接用于演示。",
     "不能虚构公告链接。演示时也应使用明确标注的 mock 数据或真实可核验 URL。"),
    ("旧知识库中的施工公告能当作今天仍然有效吗？",
     "不能。需要核验公告有效期和最新发布信息，过期公告不得当作当前事实。"),
    ("公告只写了封闭道路，能帮我补充绕行方案吗？",
     "不能补充公告未提供的绕行路线。可以建议用户结合实时导航选择合规方案。"),
    ("没有来源 URL 的交通管制消息可以直接回答用户吗？",
     "不应作为已核验事实回答。需要补充正式来源，或明确说明消息暂不可验证。"),
    ("怎样回答无法检索到的成都交管公告？",
     "应明确说明当前没有检索到可核验公告，建议查看成都公安交管官方发布渠道，不得编造内容。"),
]


def from_guardrail_boundaries() -> list[dict]:
    rows = [
        sample(
            prompt, output, category="parking_realtime_boundary",
            source="synthetic:parking_realtime_boundary",
            group_id=f"parking_realtime_boundary:{index}",
        )
        for index, (prompt, output) in enumerate(REALTIME_PARKING_BOUNDARIES)
    ]
    rows.extend(
        sample(
            prompt, output, category="rag_grounding_boundary",
            source="synthetic:rag_grounding_boundary",
            group_id=f"rag_grounding_boundary:{index}",
        )
        for index, (prompt, output) in enumerate(RAG_GROUNDING_BOUNDARIES)
    )
    return rows


def from_unavailable_tools() -> list[dict]:
    cases = [
        ("天气工具未返回数据", "weather_query",
         "天气工具当前未返回可用结果，无法确认实时天气。请稍后重试或查看官方天气服务。"),
        ("天气工具缺少访问凭据", "weather_query",
         "天气服务尚未配置访问凭据，当前不能查询实时天气。请完成配置后重试。"),
        ("路径规划工具未返回数据", "route_plan",
         "路径规划工具当前未返回可用结果，无法编造路线。请补充起终点后重试或使用实时导航。"),
        ("路径规划工具查询超时", "route_plan",
         "路径规划服务本次查询超时，暂时无法确认道路和耗时。请稍后重试或使用实时导航。"),
        ("交通流预测结果文件不存在", "ForecastTool.summarize",
         "当前没有可用的交通流预测结果文件，无法给出预测指标或趋势结论。"),
        ("交通流预测结果文件无法解析", "ForecastTool.summarize",
         "交通流预测结果文件当前无法解析，不能据此生成指标结论。请先检查实验产物。"),
        ("交通标志检测结果文件不存在", "DetectionTool.summarize",
         "当前没有可用的检测结果文件，无法确认识别类别或模型精度。"),
        ("交通标志检测结果为空", "DetectionTool.summarize",
         "检测工具没有返回有效目标，当前无法确认交通标志类别。请结合原图和模型日志复核。"),
    ]
    return [
        sample(
            "请根据工具状态给出简洁、可核验的兜底回复。",
            output, input_text=f"工具状态：{status}\n工具：{tool}",
            category="tool_unavailable", source=f"tool:{tool}",
            group_id=f"tool_unavailable:{tool}:{status}",
        )
        for status, tool, output in cases
    ]


def deduplicate(samples: list[dict]) -> list[dict]:
    seen, result = set(), []
    for row in samples:
        key = (row["instruction"], row["input"], row["output"])
        if key not in seen:
            seen.add(key)
            result.append(row)
    return result


def validate_samples(samples: list[dict]) -> dict:
    """Fail fast on patterns that would teach fabricated or placeholder facts."""
    issues = []
    sensitive = {
        "route", "weather", "tool_call", "refusal", "tool_unavailable",
        "parking_realtime_boundary", "rag_grounding_boundary",
    }
    output_counts = Counter(row["output"] for row in samples)
    for index, row in enumerate(samples):
        blob = "\n".join((row["instruction"], row["input"], row["output"], row["source"]))
        if re.search(r"\bNone\b", blob):
            issues.append(f"sample {index}: contains literal None")
        if "/example/" in blob:
            issues.append(f"sample {index}: contains placeholder /example/ URL")
        if row["category"] == "bus" and "定班制" in blob:
            issues.append(f"sample {index}: bus answer contains unspecified fixed-schedule data")
        if row["category"] in {"route", "weather"} and not row["input"]:
            issues.append(f"sample {index}: {row['category']} answer lacks tool-result input")
        if row["category"] in sensitive and output_counts[row["output"]] > 8:
            issues.append(
                f"sample {index}: {row['category']} output repeats "
                f"{output_counts[row['output']]} times"
            )
    if issues:
        preview = "\n".join(f"- {issue}" for issue in issues[:12])
        raise ValueError(f"dataset quality validation failed:\n{preview}")
    return {
        "max_output_repeat": max(output_counts.values(), default=0),
        "sensitive_max_output_repeat": max(
            (output_counts[row["output"]] for row in samples if row["category"] in sensitive),
            default=0,
        ),
    }


def split(samples: list[dict]) -> tuple[list[dict], list[dict]]:
    """Deterministic 90/10 split with semantic-group isolation."""
    rng = random.Random(RANDOM_SEED)
    groups: dict[str, list[dict]] = defaultdict(list)
    for row in samples:
        groups[row["group_id"]].append(row)
    by_category: dict[str, list[list[dict]]] = defaultdict(list)
    for rows in groups.values():
        by_category[rows[0]["category"]].append(rows)
    train, val = [], []
    for category in sorted(by_category):
        category_groups = by_category[category]
        rng.shuffle(category_groups)
        val_count = 1 if len(category_groups) >= 5 else 0
        val_count = max(val_count, round(len(category_groups) * 0.1))
        for rows in category_groups[:val_count]:
            val.extend(rows)
        for rows in category_groups[val_count:]:
            train.extend(rows)
    rng.shuffle(train)
    rng.shuffle(val)
    return train, val


def write_jsonl(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> None:
    cfg = Config.load()
    raw_dir = cfg.resolve_path(cfg.get("paths", "raw_dir", default="data/raw"))
    docs = load_raw(raw_dir) or scrape(raw_dir, offline=True)

    samples = deduplicate(
        from_knowledge(docs)
        + from_bus_stops()
        + from_bus_lines()
        + from_metro()
        + from_parking()
        + from_toll_stations()
        + from_rail_transit()
        + from_catalog()
        + from_handcrafted()
        + from_tool_calls()
        + from_weather_tools()
        + from_route_tools()
        + from_refusals()
        + from_guardrail_boundaries()
        + from_unavailable_tools()
    )
    quality = validate_samples(samples)
    train, val = split(samples)
    xindu_rows = read_csv("xindu_bus_lines.csv")
    catalog_rows = read_catalog_zip("chengdu_open_catalog_2026-06-01.zip")
    catalog_samples = [row for row in samples if row["category"] == "open_catalog"]
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    write_jsonl(OUT, samples)
    write_jsonl(TRAIN_OUT, train)
    write_jsonl(VAL_OUT, val)

    stats = {
        "seed": RANDOM_SEED,
        "total": len(samples),
        "train": len(train),
        "validation": len(val),
        "groups": len({row["group_id"] for row in samples}),
        "train_groups": len({row["group_id"] for row in train}),
        "validation_groups": len({row["group_id"] for row in val}),
        "group_leakage": len(
            {row["group_id"] for row in train} & {row["group_id"] for row in val}
        ),
        "quality": quality,
        "source_audit": {
            "xindu_bus_lines": {
                "rows": len(xindu_rows),
                "fixed_schedule_rows_excluded": sum(
                    row["首班时间"] == "定班制" or row["末班时间"] == "定班制"
                    for row in xindu_rows
                ),
            },
            "open_catalog": {
                "rows": len(catalog_rows),
                "departments": len({row["部门名称"] for row in catalog_rows}),
                "rows_with_known_file_formats": sum(
                    bool(catalog_formats(row.get("文件资源格式", "")))
                    for row in catalog_rows
                ),
                "selected_rows": len(catalog_samples),
                "selected_departments": len({
                    row["department"] for row in catalog_samples
                }),
                "selected_api_interfaces": sum(
                    row["access_method"] == "API接口" for row in catalog_samples
                ),
                "selected_rows_with_known_file_formats": sum(
                    bool(row["file_formats"]) for row in catalog_samples
                ),
            },
        },
        "categories": dict(sorted(
            (category, sum(row["category"] == category for row in samples))
            for category in {row["category"] for row in samples}
        )),
        "outputs": [OUT.name, TRAIN_OUT.name, VAL_OUT.name],
    }
    STATS_OUT.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[build_dataset] 写出 {len(samples)} 条 SFT 样本 -> {OUT}")
    print(f"[build_dataset] 划分 train={len(train)} / val={len(val)} -> {DATA_DIR}")


if __name__ == "__main__":
    main()

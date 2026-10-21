from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RESULTS_DIR = Path(r"D:/大三下/智能交通/大作业1/results")
OUTPUT_PATH = RESULTS_DIR / "paper_tables.xlsx"

DATASETS = ["METRLA", "PEMS03", "PEMS04", "PEMS08", "PEMSBAY"]
MODEL_ORDER = ["baseline", "dgl", "la", "mst"]
MODEL_LABELS = {
    "baseline": "Baseline",
    "dgl": "DGL",
    "la": "LA",
    "mst": "MST",
}

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN_SIDE = Side(style="thin", color="BFBFBF")
ALL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
CENTER = Alignment(horizontal="center", vertical="center")


def try_float(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def format_metric(value: float | None, digits: int = 4) -> str:
    if value is None:
        return "—"
    return f"{value:.{digits}f}"


def format_params(value: float | None) -> str:
    if value is None:
        return "—"
    if value >= 1_000_000:
        return f"{value / 1_000_000:.2f}M"
    if value >= 1_000:
        return f"{value / 1_000:.2f}K"
    return f"{value:.0f}"


def read_json(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def parse_log_metadata(path: Path, model_label: str) -> dict[str, Any]:
    if not path.exists():
        return {}

    text = path.read_text(encoding="utf-8", errors="ignore")
    info: dict[str, Any] = {
        "source": path.name,
        "model_name": model_label,
    }

    param_match = re.search(r"Total parameters:\s*([\d,]+)", text)
    if param_match:
        info["total_params"] = try_float(param_match.group(1))

    trainset_match = re.search(r"Trainset:\s*x-\((\d+),\s*(\d+),\s*(\d+),\s*(\d+)\)", text)
    if trainset_match:
        info["train_samples"] = int(trainset_match.group(1))
        info["in_steps"] = int(trainset_match.group(2))
        info["num_nodes"] = int(trainset_match.group(3))
        info["input_channels"] = int(trainset_match.group(4))

    val_match = re.search(r"Valset:\s*x-\((\d+),", text)
    if val_match:
        info["val_samples"] = int(val_match.group(1))

    test_match = re.search(r"Testset:\s*x-\((\d+),", text)
    if test_match:
        info["test_samples"] = int(test_match.group(1))

    model_match = re.search(r"Model:\s*(.+)", text)
    if model_match:
        info["model_name"] = model_match.group(1).strip()

    metric_match = re.search(
        r"All Steps RMSE\s*=\s*([\d.]+),\s*MAE\s*=\s*([\d.]+),\s*MAPE\s*=\s*([\d.]+)",
        text,
    )
    if metric_match:
        info["rmse"] = try_float(metric_match.group(1))
        info["mae"] = try_float(metric_match.group(2))
        info["mape"] = try_float(metric_match.group(3))

    time_match = re.search(r"Inference time:\s*([\d.]+)\s*s", text)
    if time_match:
        info["inference_time"] = try_float(time_match.group(1))

    epoch_match = re.search(r'"best_epoch"\s*:\s*(\d+)', text)
    if epoch_match:
        info["best_epoch"] = int(epoch_match.group(1))

    if metric_match:
        info["status"] = "completed_from_log"
    elif info.get("train_samples") is not None or info.get("total_params") is not None:
        info["status"] = "log_only"
    else:
        info["status"] = "missing"
    return info


def parse_eff_log(path: Path) -> dict[str, Any]:
    info = parse_log_metadata(path, "STAEformer-EFF")
    if not info:
        return {}
    if info.get("status") == "log_only":
        info["status"] = "final_metrics_missing"
    return info


def collect_results() -> dict[str, dict[str, dict[str, Any]]]:
    all_results: dict[str, dict[str, dict[str, Any]]] = {}
    for dataset in DATASETS:
        dataset_results: dict[str, dict[str, Any]] = {}
        for model in MODEL_ORDER:
            if model == "eff":
                eff_payload = parse_eff_log(RESULTS_DIR / f"eff_{dataset}_train.log")
                if eff_payload:
                    dataset_results[model] = eff_payload
                continue

            json_path = RESULTS_DIR / f"{model}_{dataset}_results.json"
            payload = read_json(json_path)
            if payload is not None:
                dataset_results[model] = payload
                continue

            log_payload = parse_log_metadata(
                RESULTS_DIR / f"{model}_{dataset}_train.log",
                MODEL_LABELS[model],
            )
            if log_payload:
                dataset_results[model] = log_payload
        all_results[dataset] = dataset_results
    return all_results


def extract_row(dataset: str, model: str, payload: dict[str, Any]) -> dict[str, Any]:
    if model == "eff" or "test_result" not in payload:
        return {
            "dataset": dataset,
            "model": MODEL_LABELS[model],
            "model_name": payload.get("model_name", MODEL_LABELS[model]),
            "mae": try_float(payload.get("mae")),
            "rmse": try_float(payload.get("rmse")),
            "mape": try_float(payload.get("mape")),
            "inference_time": try_float(payload.get("inference_time")),
            "total_params": try_float(payload.get("total_params")),
            "best_epoch": payload.get("best_epoch"),
            "train_samples": payload.get("train_samples"),
            "val_samples": payload.get("val_samples"),
            "test_samples": payload.get("test_samples"),
            "num_nodes": payload.get("num_nodes"),
            "status": payload.get("status", "missing"),
            "source": payload.get("source", "—"),
        }

    train_result = payload.get("train_result", {})
    test_result = payload.get("test_result", {})
    model_args = payload.get("model_args", {})
    return {
        "dataset": dataset,
        "model": MODEL_LABELS[model],
        "model_name": payload.get("model_name", MODEL_LABELS[model]),
        "mae": try_float(test_result.get("mae")),
        "rmse": try_float(test_result.get("rmse")),
        "mape": try_float(test_result.get("mape")),
        "inference_time": try_float(test_result.get("inference_time")),
        "total_params": try_float(payload.get("total_params")),
        "best_epoch": train_result.get("best_epoch"),
        "train_samples": None,
        "val_samples": None,
        "test_samples": None,
        "num_nodes": model_args.get("num_nodes"),
        "status": "completed",
        "source": f"{model}_{dataset}_results.json",
    }


def build_rows(all_results: dict[str, dict[str, dict[str, Any]]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for dataset in DATASETS:
        for model in MODEL_ORDER:
            payload = all_results.get(dataset, {}).get(model)
            if payload is None:
                rows.append(
                    {
                        "dataset": dataset,
                        "model": MODEL_LABELS[model],
                        "model_name": MODEL_LABELS[model],
                        "mae": None,
                        "rmse": None,
                        "mape": None,
                        "inference_time": None,
                        "total_params": None,
                        "best_epoch": None,
                        "train_samples": None,
                        "val_samples": None,
                        "test_samples": None,
                        "num_nodes": None,
                        "status": "missing",
                        "source": "—",
                    }
                )
                continue
            rows.append(extract_row(dataset, model, payload))
    return rows


def apply_best_highlight(ws: Worksheet, row_index: int, column_count: int) -> None:
    for col in range(1, column_count + 1):
        ws.cell(row_index, col).fill = HIGHLIGHT_FILL


def style_table(ws: Worksheet, max_col: int | None = None) -> None:
    max_col = max_col or ws.max_column
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = ALL_BORDER
            cell.alignment = CENTER
            if cell.row == 1:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL

    for column_index in range(1, max_col + 1):
        max_length = 0
        for row_index in range(1, ws.max_row + 1):
            value = ws.cell(row_index, column_index).value
            cell_text = "" if value is None else str(value)
            max_length = max(max_length, len(cell_text))
        ws.column_dimensions[get_column_letter(column_index)].width = min(max_length + 4, 28)

    ws.freeze_panes = "A2"


def build_summary_notes(rows: list[dict[str, Any]]) -> list[str]:
    notes = ["1. 黄色行为该数据集 RMSE 最优模型。"]

    pems07_partial = [
        row["model"]
        for row in rows
        if row["dataset"] == "PEMS07" and row["status"] in {"log_only", "final_metrics_missing"}
    ]
    if pems07_partial:
        notes.append(f"2. PEMS07 的 {', '.join(pems07_partial)} 仅检测到日志/元数据，缺少最终测试指标。")

    eff_partial = [row["dataset"] for row in rows if row["model"] == "EFF" and row["status"] == "final_metrics_missing"]
    if eff_partial:
        notes.append(f"3. EFF 在 {', '.join(eff_partial)} 仅有参数与样本规模信息，未检测到最终测试指标。")

    missing_rows = [
        f"{row['dataset']}-{row['model']}"
        for row in rows
        if row["status"] == "missing"
    ]
    if missing_rows:
        notes.append(f"4. 以下条目未发现结果文件或有效日志：{', '.join(missing_rows)}。")

    notes.append("5. 表格可直接复制到论文，建议指标保留四位小数。")
    return notes

def add_summary_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "总表"
    headers = ["数据集", "模型", "MAE", "RMSE", "MAPE(%)", "推理时间(s)", "参数量", "状态", "数据源"]
    ws.append(headers)

    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(row["dataset"], []).append(row)

    for dataset in DATASETS:
        dataset_rows = grouped.get(dataset, [])
        completed = [r for r in dataset_rows if r["rmse"] is not None]
        best_rmse = min((r["rmse"] for r in completed), default=None)
        for row in dataset_rows:
            ws.append(
                [
                    row["dataset"],
                    row["model"],
                    row["mae"],
                    row["rmse"],
                    row["mape"],
                    row["inference_time"],
                    format_params(row["total_params"]),
                    row["status"],
                    row["source"],
                ]
            )
            current_row = ws.max_row
            if row["mae"] is not None:
                ws.cell(current_row, 3).number_format = "0.0000"
            if row["rmse"] is not None:
                ws.cell(current_row, 4).number_format = "0.0000"
            if row["mape"] is not None:
                ws.cell(current_row, 5).number_format = "0.0000"
            if row["inference_time"] is not None:
                ws.cell(current_row, 6).number_format = "0.00"
            if best_rmse is not None and row["rmse"] == best_rmse:
                apply_best_highlight(ws, current_row, len(headers))
        ws.append([""] * len(headers))

    notes = build_summary_notes(rows)
    ws["K1"] = "说明"
    ws["K1"].font = Font(bold=True)
    for index, note in enumerate(notes, start=2):
        ws[f"K{index}"] = note

    style_table(ws, max_col=11)


def add_dataset_sheet(wb: Workbook, dataset: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(dataset)
    headers = ["模型", "完整模型名", "MAE", "RMSE", "MAPE(%)", "推理时间(s)", "参数量", "最佳轮次", "状态"]
    ws.append(headers)

    dataset_rows = [row for row in rows if row["dataset"] == dataset]
    completed = [r for r in dataset_rows if r["rmse"] is not None]
    best_rmse = min((r["rmse"] for r in completed), default=None)

    for row in dataset_rows:
        ws.append(
            [
                row["model"],
                row["model_name"],
                row["mae"],
                row["rmse"],
                row["mape"],
                row["inference_time"],
                format_params(row["total_params"]),
                row["best_epoch"] if row["best_epoch"] is not None else "—",
                row["status"],
            ]
        )
        current_row = ws.max_row
        if row["mae"] is not None:
            ws.cell(current_row, 3).number_format = "0.0000"
        if row["rmse"] is not None:
            ws.cell(current_row, 4).number_format = "0.0000"
        if row["mape"] is not None:
            ws.cell(current_row, 5).number_format = "0.0000"
        if row["inference_time"] is not None:
            ws.cell(current_row, 6).number_format = "0.00"
        if best_rmse is not None and row["rmse"] == best_rmse:
            apply_best_highlight(ws, current_row, len(headers))

    ws.append([])
    ws.append(["论文表述建议", "按 RMSE、MAE、MAPE 三项指标比较各模型；参数量和推理时间用于效率分析。"])
    style_table(ws, max_col=len(headers))


def add_dataset_info_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("数据集信息")
    headers = ["数据集", "模型", "节点数", "训练样本", "验证样本", "测试样本", "参数量", "状态", "来源"]
    ws.append(headers)

    for row in rows:
        ws.append(
            [
                row["dataset"],
                row["model"],
                row["num_nodes"] if row["num_nodes"] is not None else "—",
                row["train_samples"] if row["train_samples"] is not None else "—",
                row["val_samples"] if row["val_samples"] is not None else "—",
                row["test_samples"] if row["test_samples"] is not None else "—",
                format_params(row["total_params"]),
                row["status"],
                row["source"],
            ]
        )

    style_table(ws)


def main() -> None:
    all_results = collect_results()
    rows = build_rows(all_results)

    workbook = Workbook()
    add_summary_sheet(workbook, rows)
    for dataset in DATASETS:
        add_dataset_sheet(workbook, dataset, rows)
    add_dataset_info_sheet(workbook, rows)
    workbook.save(OUTPUT_PATH)
    print(f"Saved xlsx: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()

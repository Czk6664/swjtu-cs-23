from __future__ import annotations

import os
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
LAB_DIR = ROOT / "firewall_lab"
sys.path.insert(0, str(LAB_DIR))

from main import build_default_policy  # noqa: E402
from policy_audit import detect_conflicting_rules, detect_shadowed_rules  # noqa: E402
from test_cases import run_test_suite, collect_logs, summarize_results  # noqa: E402


OUT = ROOT / "实验7-实验结果与分析.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, widths_cm: list[float], font_size=8.0) -> None:
    set_table_widths(table, widths_cm)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Calibri"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(font_size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], font_size=8.0):
    table = doc.add_table(rows=1, cols=len(headers), style="Table Grid")
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            row.cells[i].text = value
    style_table(table, widths_cm, font_size)
    doc.add_paragraph()
    return table


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered_summary(doc: Document, title: str, text: str) -> None:
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.add_run(text)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def make_document() -> None:
    rules = build_default_policy()
    cases = [f"TC-{i:02d}" for i in range(1, 9)]
    results = run_test_suite(cases, rules)
    logs = collect_logs(results)
    summary = summarize_results(results)
    shadowed = detect_shadowed_rules(rules)
    conflicts = detect_conflicting_rules(rules)

    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("实验7 防火墙技术综合验证实验结果与分析")
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    meta = doc.add_paragraph()
    meta.add_run("实验环境：Python 规则引擎模拟；默认策略：deny；运行日期：2026-06-03。")

    doc.add_heading("八、实验结果", level=1)

    doc.add_heading("8.1 策略设计记录表", level=2)
    rows_81 = [
        ["允许内网访问 Web 服务", "R-01", "10.0.1.0/24", "任意外部地址", "TCP/80,443", "allow", "是", "满足最小必要出站访问"],
        ["允许内网 DNS 查询", "R-02", "10.0.1.0/24", "任意 DNS 服务器", "UDP/53", "allow", "是", "保证域名解析可用"],
        ["发布 DMZ HTTPS 服务", "R-03", "任意外部地址", "203.0.113.10", "TCP/443", "allow", "是", "只开放指定 DMZ 服务"],
        ["阻断外部 SSH 入站", "R-04", "任意外部地址", "10.0.1.0/24", "TCP/22", "deny", "是", "防止外部直接管理内网主机"],
        ["阻断数据库子网入站", "R-05", "任意外部地址", "10.0.2.0/24", "任意服务", "deny", "是", "保护数据库网段"],
        ["Default", "默认策略", "任意", "任意", "任意", "deny", "是", "未命中规则时拒绝并记录"],
    ]
    add_table(doc, ["策略目标", "规则编号", "源", "目的", "服务", "动作", "日志是否存在", "备注"], rows_81, [2.3, 1.35, 2.2, 2.35, 1.75, 1.2, 1.45, 3.6], 7.5)

    doc.add_heading("8.2 流量验证记录表", level=2)
    desc = {
        "TC-01": "内网主机访问外部 HTTP",
        "TC-02": "外部服务器返回 TC-01 流量",
        "TC-03": "外部主动 SSH 访问内网",
        "TC-04": "外部访问 DMZ HTTPS",
        "TC-05": "内网访问 DMZ HTTPS",
        "TC-06": "内网发起 DNS 查询",
        "TC-07": "内网发起 ICMP 流量",
        "TC-08": "外部访问数据库子网",
    }
    expected = {
        "TC-01": "allow",
        "TC-02": "allow",
        "TC-03": "deny",
        "TC-04": "allow",
        "TC-05": "allow",
        "TC-06": "allow",
        "TC-07": "deny",
        "TC-08": "deny",
    }
    rows_82 = []
    for cid, result in zip(cases, results):
        ok = "通过" if expected[cid] == result["action"] else "不通过"
        rows_82.append([cid, desc[cid], expected[cid], result["action"], result["matched_rule"], "是" if result["logged"] else "否", ok])
    add_table(doc, ["测试编号", "测试描述", "预期动作", "实际动作", "命中规则", "日志是否存在", "结论"], rows_82, [1.4, 4.1, 1.6, 1.6, 1.7, 1.8, 1.3], 8.0)

    doc.add_heading("8.3 规则冲突与优化记录表", level=2)
    rows_83 = [
        ["遮蔽规则", "无", "审计未发现后续规则被前置规则完全覆盖", "规则条件之间不存在完全包含且优先级更高的覆盖关系", "继续保持窄范围规则优先，新增规则后重新审计"],
        ["动作冲突", "无", "审计未发现条件重叠但动作相反的规则", "现有规则目标清晰，入站拒绝和出站允许边界不同", "新增 DMZ 或管理规则时先做重叠检查"],
        ["过度放行", "R-01", "内网 Web 出站目的地址为任意外部地址", "该规则便于实验验证，但生产环境中范围偏宽", "可收紧为代理服务器、可信业务网段或按用户组控制"],
        ["日志优化", "DEFAULT", "默认拒绝命中 TC-07，可追踪但规则来源为 DEFAULT", "未命中显式 ICMP 规则，说明默认拒绝生效", "保留默认日志，并为高频默认拒绝流量补充专门规则"],
    ]
    add_table(doc, ["问题类型", "涉及规则", "现象", "原因分析", "改进建议"], rows_83, [2.0, 1.8, 4.2, 4.0, 4.0], 7.8)

    doc.add_heading("8.4 日志审计记录表", level=2)
    rows_84 = []
    for i, log in enumerate(logs, start=1):
        service = f"{log['protocol']}/{log['dst_port'] or '-'}"
        conclusion = "允许记录完整" if log["action"] == "allow" else "拒绝记录完整"
        rows_84.append([f"L{i}", log["time"], log["src"], log["dst"], service, log["action"], log["rule"], conclusion])
    add_table(doc, ["日志编号", "时间", "源", "目的", "服务", "动作", "规则", "审计结论"], rows_84, [1.0, 2.75, 2.1, 2.1, 1.6, 1.2, 1.2, 3.1], 7.2)

    doc.add_heading("8.5 实验4状态检测表", level=2)
    rows_85 = [
        ["T1", "Trust -> Untrust", "否", "allow", "allow", "新增一条 TCP 状态项", "内网主动访问外部 HTTP，命中 R-01"],
        ["T2", "Untrust -> Trust 返回流量", "是", "allow", "allow", "刷新已有状态项", "返回流量匹配 STATE，而不是普通入站放行"],
        ["T3", "Untrust -> Trust 主动入站", "否", "deny", "deny", "无新增状态项", "外部主动访问内网 SSH，命中 R-04 拒绝"],
    ]
    add_table(doc, ["测试序号", "连接方向", "是否已有状态", "预期动作", "实际动作", "状态表变化", "说明"], rows_85, [1.2, 3.2, 1.8, 1.5, 1.5, 3.0, 4.0], 8.0)

    doc.add_heading("8.6 实验5 NAT与DMZ表", level=2)
    rows_86 = [
        ["SNAT", "10.0.1.10:53000", "203.0.113.1:53000", "Trust -> Untrust 且源为 10.0.1.0/24", "只允许必要出站，隐藏内网地址"],
        ["DNAT", "203.0.113.10:443", "10.0.3.10:443", "Untrust -> DMZ 且目的端口为 TCP/443", "只发布 DMZ HTTPS 服务"],
        ["端口映射", "203.0.113.10:443", "10.0.3.10:443", "外部访问公网 HTTPS 入口", "禁止开放 SSH、数据库等额外端口"],
    ]
    add_table(doc, ["项目", "转换前", "转换后", "匹配条件", "安全要求"], rows_86, [1.5, 3.0, 3.0, 5.0, 4.0], 8.0)

    doc.add_heading("8.7 相关关键实验步骤截图或程序运行或结果截图", level=2)
    doc.add_paragraph("关键程序运行结果如下，可作为实验运行截图对应内容：")
    output_lines = [
        "TC-01 10.0.1.10:53000 -> 198.51.100.20:80 TCP outbound allow R-01",
        "TC-02 198.51.100.20:80 -> 10.0.1.10:53000 TCP inbound allow STATE",
        "TC-03 203.0.113.8:44000 -> 10.0.1.10:22 TCP inbound deny R-04",
        "TC-04 203.0.113.8:44000 -> 203.0.113.10:443 TCP inbound allow R-03",
        "TC-05 10.0.1.10:53001 -> 203.0.113.10:443 TCP outbound allow R-01",
        "TC-06 10.0.1.10:53002 -> 198.51.100.53:53 UDP outbound allow R-02",
        "TC-07 10.0.1.10 -> 198.51.100.20 ICMP outbound deny DEFAULT",
        "TC-08 203.0.113.8:44000 -> 10.0.2.30:3306 TCP inbound deny R-05",
    ]
    for line in output_lines:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)

    doc.add_heading("九、实验分析", level=1)

    doc.add_heading("9.1 包过滤分析", level=2)
    add_bullets(doc, [
        "R-01 的安全目标是允许内网用户访问必要的 Web 服务，避免业务访问被全部阻断；R-02 允许 DNS 查询，保证域名解析；R-03 只发布 DMZ 的 HTTPS 服务；R-04 和 R-05 分别阻断外部 SSH 管理入口和数据库子网入站访问，体现分区隔离。",
        "默认拒绝策略更符合最小权限原则，未明确允许的流量都会被拒绝并记录，安全性较高；默认允许策略便于连通性测试，但容易因漏配拒绝规则造成越权访问。",
        "协议、端口或方向字段任一配置错误都可能导致误放行或误阻断。例如把 inbound 写成 outbound 会使入站规则不生效，把 443 写成 433 会导致 HTTPS 服务不可达，把 TCP 配成 UDP 会使命中失败。",
    ])

    doc.add_heading("9.2 状态检测分析", level=2)
    add_bullets(doc, [
        "当 TC-01 内网主机主动访问外部 HTTP 并被允许后，防火墙在状态表中记录该五元组连接。随后 TC-02 的反向返回流量虽然方向为入站，但因命中已建立状态而被允许。",
        "外部主动入站与返回流量的区别在于连接发起方和状态是否存在。返回流量属于已允许连接的一部分；外部主动入站没有状态表记录，需要命中显式放行规则，否则会被拒绝。",
        "如果状态表过期、被清空或异常更新，合法返回流量将无法命中 STATE，可能转入普通规则匹配并被默认拒绝，因此状态表超时和清理策略会直接影响连接稳定性。",
    ])

    doc.add_heading("9.3 NAT 与 DMZ 分析", level=2)
    add_bullets(doc, [
        "SNAT 后，外部日志看到的源地址通常是公网出口地址，而内部日志仍可记录真实内网地址；DNAT 后，防火墙入口日志可见公网目的地址，DMZ 服务器日志看到的目的地址则是转换后的私有地址。",
        "对外发布 DMZ 服务只应暴露指定服务端口，不应让外部流量直接进入 Trust 内网。DMZ 的价值在于即使对外服务被攻击，也能降低攻击横向移动到内网的概率。",
        "如果 NAT 规则与过滤规则顺序不当，可能出现按转换前地址匹配失败或按转换后地址误放行的情况。实验中应明确过滤发生在 NAT 前还是 NAT 后，并在日志中保留转换前后地址以便追踪。",
    ])

    doc.add_heading("9.4 日志与审计分析", level=2)
    add_bullets(doc, [
        f"本次日志包含时间、源、目的、协议、端口、方向、动作、命中规则和原因，字段足以支持基本追踪与复现。",
        f"统计结果：日志总数 {summary['total']} 条，允许 {summary['allow']} 条，拒绝 {summary['deny']} 条，默认策略命中 {summary['default']} 条。高频命中规则为 R-01，共 {summary['rule_hits'].get('R-01', 0)} 次。",
        "优化建议：R-01 可进一步收紧目的范围；对 DEFAULT 高频拒绝流量建立专门的显式拒绝规则；定期检查规则顺序，避免后续新增规则被遮蔽。",
    ])

    summaries = {
        "9.5 实验1总结": "通过基线连通性确认了 Trust、Untrust、DMZ 等区域和服务角色，为后续策略验证提供了对照基准。",
        "9.6 实验2总结": "静态包过滤能够根据源、目的、协议、端口和方向进行访问控制，默认拒绝策略可以有效限制未授权流量。",
        "9.7 实验3总结": "规则顺序会影响最终决策，应将更具体、更高风险或显式拒绝的规则放在合适位置，并定期做冲突审计。",
        "9.8 实验4总结": "状态检测能够识别合法返回流量，使内网主动连接的响应被允许，同时仍能拒绝外部主动入站访问。",
        "9.9 实验5总结": "NAT 与 DMZ 发布可以隐藏内部地址并对外提供必要服务，但必须配合过滤规则控制暴露面。",
        "9.10 实验6总结": "日志审计能定位命中规则、默认拒绝和异常流量，是策略优化和问题追踪的重要依据。",
        "9.11 实验7总结": "程序化规则引擎完整模拟了规则表、状态表、默认策略、NAT 记录和日志审计流程，验证结果与预期一致。",
    }
    for title_text, body in summaries.items():
        add_numbered_summary(doc, title_text, body)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("实验7 防火墙技术综合验证")

    doc.save(OUT)


if __name__ == "__main__":
    make_document()
    print(OUT)

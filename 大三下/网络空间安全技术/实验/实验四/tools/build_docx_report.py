from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from otp_lab import main as lab


OUTPUT_PATH = PROJECT_ROOT / "otp_lab_report.docx"


def set_east_asia_font(run, font_name: str = "SimSun") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 8) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(text)) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    set_east_asia_font(run)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_paragraph(document: Document, text: str, size: int = 10.5, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_east_asia_font(run)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(31, 78, 121)
    set_east_asia_font(run, "SimHei")


def add_table(document: Document, headers: list[str], rows: list[dict[str, str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "D9EAF7")
        set_cell_text(cell, header, bold=True, size=8)

    for item in rows:
        cells = table.add_row().cells
        for index, header in enumerate(headers):
            set_cell_text(cells[index], item.get(header, ""), size=7)

    document.add_paragraph()


def build_docx() -> Path:
    lab.LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    lab.LOG_PATH.write_text("", encoding="utf-8")
    static_rows = lab.run_static_cases()
    list_rows = lab.run_list_cases()
    counter_rows = lab.run_counter_cases()
    time_rows = lab.run_time_cases()
    security_rows, comparison_rows, fail_summary = lab.run_security_cases()
    lab.REPORT_PATH.write_text(
        lab.build_report(
            static_rows,
            list_rows,
            counter_rows,
            time_rows,
            security_rows,
            comparison_rows,
            fail_summary,
        ),
        encoding="utf-8-sig",
    )

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    styles = document.styles
    styles["Normal"].font.name = "SimSun"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    styles["Normal"].font.size = Pt(10.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("一次性口令机制（OTP）验证实验报告")
    run.bold = True
    run.font.size = Pt(20)
    set_east_asia_font(run, "SimHei")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run("姓名：__________    学号：__________    班级：__________")
    meta_run.font.size = Pt(10.5)
    set_east_asia_font(meta_run)

    add_heading(document, "一、实验名称", 1)
    add_paragraph(document, "一次性口令机制（OTP）验证实验。")

    add_heading(document, "二、实验目的", 1)
    for item in [
        "理解静态口令与一次性口令在认证安全性上的差异，重点说明重放攻击为什么会威胁静态口令。",
        "掌握共享口令表型、计数器/序列同步型、时间同步型三类 OTP 的生成与验证方式。",
        "实现认证端生成 OTP、验证端重算比较、状态推进、复用检测和日志记录。",
        "验证 OTP 对重放攻击、字段篡改、计数器漂移和时钟漂移的抵抗能力。",
        "总结 OTP 的安全边界：它增强认证与新鲜性，不直接等价于通信保密。",
    ]:
        add_paragraph(document, item)

    add_heading(document, "三、实验原理", 1)
    add_paragraph(
        document,
        "静态口令每次认证都提交同一秘密值，攻击者截获一次合法消息后可以原样重放。OTP 将口令与当前位置、计数器或时间槽绑定，服务器成功验证后推进状态或记录已用标记，使旧消息再次提交时被识别为复用或旧状态。",
    )
    add_paragraph(
        document,
        "共享口令表型依赖双方一致的 pos；计数器型使用 HMAC-SHA256(secret, counter) 派生 OTP，并可用 look_ahead 滑动窗口容忍小幅超前；时间同步型按固定 step 划分时间槽，验证端在 allowed_window 内容忍时钟偏差，同时记录已用 slot。",
    )

    add_heading(document, "四、实验环境", 1)
    add_paragraph(document, "操作系统：Windows / Linux / macOS 均可；本次在 Windows PowerShell 环境运行。")
    add_paragraph(document, "编程语言：Python 3.13.5。")
    add_paragraph(document, "开发工具：命令行 + 模块化 Python 程序。")

    add_heading(document, "五、程序设计", 1)
    add_paragraph(
        document,
        "目录结构为 otp_lab/common.py、static_password_baseline.py、otp_list_mode.py、otp_counter_mode.py、otp_time_mode.py、attack_simulation.py、main.py。公共模块负责规范化、HMAC 派生、截断、恒定时间比较、复用标记和日志；模式模块负责生成、验证和状态推进；攻击模块负责重放、篡改和漂移模拟。",
    )
    add_paragraph(
        document,
        f"公共参数：用户 u01/u02；口令表 L=[111111,222222,333333,444444,555555]；计数器初始值 0，look_ahead={lab.LOOK_AHEAD}；时间模式 step={lab.STEP}s，allowed_window={lab.ALLOWED_WINDOW}，t0={lab.BASE_TS}。",
    )

    detail_headers = ["用例编号", "操作描述", "测试输入", "预期结果", "实际输出与分析", "结果判定"]
    add_heading(document, "六、实验结果", 1)
    add_heading(document, "实验 A：静态口令对照", 2)
    add_table(document, detail_headers, static_rows)
    add_heading(document, "实验 B：共享口令表型 OTP", 2)
    add_table(document, detail_headers, list_rows)
    add_heading(document, "实验 C：计数器/序列同步型 OTP", 2)
    add_table(document, detail_headers, counter_rows)
    add_heading(document, "实验 D：时间同步型 OTP", 2)
    add_table(document, detail_headers, time_rows)
    add_heading(document, "实验 E/F：统一安全性用例与模式比较", 2)
    add_table(document, ["用例编号", "攻击类型", "测试说明", "口令表", "计数器", "时间 OTP", "原因"], security_rows)
    add_table(document, ["安全属性", "口令表", "计数器", "时间 OTP"], comparison_rows)

    add_heading(document, "七、结果分析", 1)
    for item in [
        "静态口令只做固定字符串比较，不维护新鲜性状态；完整重放仍可通过。",
        "共享口令表型 OTP 成功后推进 pos 并记录已用标记，旧位置口令或完整重放会被拒绝。",
        "计数器型 OTP 将 OTP 与 counter 绑定；look_ahead 内漂移可通过并推进服务器状态，窗口外漂移应重同步或重新注册。",
        "时间同步型 OTP 将 OTP 与时间槽绑定；允许窗口可容忍小幅时钟误差，但已用 slot 仍必须拒绝重放。",
        "OTP 提升认证新鲜性和抗重放能力，不提供通信保密，真实系统仍需 TLS、消息完整性保护、限速和审计。",
    ]:
        add_paragraph(document, item)

    add_heading(document, "八、运行方法", 1)
    add_paragraph(document, "在实验目录执行：python -m otp_lab.main。运行后会生成 otp_lab/logs/events.jsonl 和 otp_lab_report.md。")

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_docx())

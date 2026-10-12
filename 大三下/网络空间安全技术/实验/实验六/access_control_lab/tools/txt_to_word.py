import builtins
import io
import keyword
import re
import token
import tokenize
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "第六部分代码填充版.txt"
OUTPUT = ROOT / "第六部分代码填充版.docx"


COLORS = {
    "default": "24292F",
    "keyword": "0000FF",
    "string": "008000",
    "comment": "808080",
    "number": "B36B00",
    "builtin": "795E26",
    "decorator": "AF00DB",
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D0D7DE", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def clear_paragraph_spacing(paragraph):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def split_sections(text):
    pattern = re.compile(r"^(6\.(?:[2-9]|10|11) .+)$", re.MULTILINE)
    matches = list(pattern.finditer(text))
    sections = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        title = match.group(1).strip()
        code = text[start:end].strip("\r\n")
        sections.append((title, code))
    return sections


def token_style(tok_type, tok_text):
    if tok_type == token.STRING:
        return "string"
    if tok_type == tokenize.COMMENT:
        return "comment"
    if tok_type == token.NUMBER:
        return "number"
    if tok_type == token.NAME:
        if tok_text in keyword.kwlist:
            return "keyword"
        if tok_text in dir(builtins):
            return "builtin"
    if tok_text == "@":
        return "decorator"
    return "default"


def highlight_segments(code):
    lines = code.splitlines()
    if not lines:
        return [[]]
    styles = [["default"] * max(len(line), 1) for line in lines]
    try:
        tokens = tokenize.generate_tokens(io.StringIO(code).readline)
        for tok in tokens:
            tok_type, tok_text, start, end, _ = tok
            style = token_style(tok_type, tok_text)
            if style == "default" or tok_type in {token.ENCODING, token.ENDMARKER, token.NEWLINE, token.NL, token.INDENT, token.DEDENT}:
                continue
            srow, scol = start
            erow, ecol = end
            for row_index in range(srow - 1, erow):
                if row_index < 0 or row_index >= len(styles):
                    continue
                line_len = len(styles[row_index])
                left = scol if row_index == srow - 1 else 0
                right = ecol if row_index == erow - 1 else line_len
                right = min(right, line_len)
                for col in range(max(left, 0), max(right, 0)):
                    styles[row_index][col] = style
    except tokenize.TokenError:
        pass

    result = []
    for line, style_line in zip(lines, styles):
        if line == "":
            result.append([(" ", "default")])
            continue
        parts = []
        current_style = style_line[0]
        buffer = []
        for char, style in zip(line, style_line):
            if style != current_style:
                parts.append(("".join(buffer), current_style))
                buffer = [char]
                current_style = style
            else:
                buffer.append(char)
        if buffer:
            parts.append(("".join(buffer), current_style))
        result.append(parts)
    return result


def add_code_block(document, code):
    lines = code.splitlines()
    segments = highlight_segments(code)
    if not lines:
        lines = [""]
    for row_index, _line in enumerate(lines):
        paragraph = document.add_paragraph()
        clear_paragraph_spacing(paragraph)
        paragraph.paragraph_format.left_indent = Cm(0.15)
        paragraph.paragraph_format.right_indent = Cm(0.15)
        set_paragraph_shading(paragraph, "F6F8FA")

        number_run = paragraph.add_run(f"{row_index + 1:>4}  ")
        number_run.font.name = "Consolas"
        number_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        number_run.font.size = Pt(8)
        number_run.font.color.rgb = RGBColor.from_string("6A737D")
        for text, style in segments[row_index] if row_index < len(segments) else [(" ", "default")]:
            run = paragraph.add_run(text)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(COLORS[style])


def configure_document(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)

    heading = styles["Heading 1"]
    heading.font.name = "微软雅黑"
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string("1F4E78")


def build_docx():
    source = INPUT.read_text(encoding="utf-8")
    sections = split_sections(source)
    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("第六部分代码填充版")
    title_run.bold = True
    title_run.font.name = "微软雅黑"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("17365D")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("代码底纹灰色，Python 语法高亮，带行号，代码字体 Consolas")
    subtitle_run.font.name = "微软雅黑"
    subtitle_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor.from_string("666666")

    for title_text, code in sections:
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(4)
        heading.add_run(title_text)
        add_code_block(document, code)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()

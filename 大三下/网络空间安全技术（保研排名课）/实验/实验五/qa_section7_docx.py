import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
DOCX = max(
    [p for p in ROOT.glob("*.docx") if not p.name.startswith("~$")],
    key=lambda p: p.stat().st_mtime,
)
DATA = json.loads((ROOT / "output" / "experiment_results.json").read_text(encoding="utf-8"))


def main() -> None:
    doc = Document(DOCX)
    issues: list[str] = []

    section = doc.sections[0]
    if section.orientation != 1:
        issues.append("page orientation is not landscape")
    if section.page_width.cm < 29 or section.page_height.cm > 22:
        issues.append("page size is not A4 landscape-like")
    usable_width = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    if usable_width < 27:
        issues.append(f"usable page width too small: {usable_width:.2f} cm")

    expected_counts = [6, 6, 7, 5, 5, 8, 8, 5]
    expected_cols = [3, 6, 6, 6, 6, 6, 6, 7]
    if len(doc.tables) != 8:
        issues.append(f"table count {len(doc.tables)} != 8")

    for idx, table in enumerate(doc.tables):
        if table.style.name != "Table Grid":
            issues.append(f"table {idx + 1} style is {table.style.name}, expected Table Grid")
        if len(table.rows) != expected_counts[idx]:
            issues.append(f"table {idx + 1} rows {len(table.rows)} != {expected_counts[idx]}")
        if len(table.columns) != expected_cols[idx]:
            issues.append(f"table {idx + 1} cols {len(table.columns)} != {expected_cols[idx]}")
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if not cell.text.strip():
                    issues.append(f"table {idx + 1} empty cell row {r + 1} col {c + 1}")
                tc_pr = cell._tc.get_or_add_tcPr()
                if tc_pr.first_child_found_in("w:tcMar") is None:
                    issues.append(f"table {idx + 1} cell row {r + 1} col {c + 1} has no cell margin")
        header_shaded = True
        for cell in table.rows[0].cells:
            shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            header_shaded = header_shaded and shd is not None and shd.get(qn("w:fill")) == "D9EAF7"
        if not header_shaded:
            issues.append(f"table {idx + 1} header shading is missing or inconsistent")

    detail_ids = [r["test_id"] for r in DATA["records"]]
    table_ids = []
    for table_idx in range(1, 7):
        table_ids.extend(row.cells[0].text.strip() for row in doc.tables[table_idx].rows[1:])
    if table_ids != detail_ids:
        issues.append(f"detail test id order/content mismatch: {table_ids} != {detail_ids}")

    cross_ids = [row.cells[0].text.strip() for row in doc.tables[7].rows[1:]]
    if cross_ids != ["X-01", "X-02", "X-03", "X-04"]:
        issues.append(f"cross test ids mismatch: {cross_ids}")

    allowed_expected = {
        "通过",
        "拒绝（比较失败）",
        "拒绝（新鲜性失败）",
        "拒绝（复用失败）",
        "拒绝（身份失败）",
        "拒绝（完整性失败）",
    }
    required_markers = ["(a) 返回值/错误码", "(b) 服务器状态变化", "(c) 关键日志字段"]
    for table_idx in range(1, 7):
        for r, row in enumerate(doc.tables[table_idx].rows[1:], start=2):
            actual = row.cells[4].text
            verdict = row.cells[5].text
            for marker in required_markers:
                if marker not in actual:
                    issues.append(f"table {table_idx + 1} row {r} missing {marker}")
            if not verdict.startswith("是。"):
                issues.append(f"table {table_idx + 1} row {r} verdict not consistent: {verdict[:40]}")
            if row.cells[3].text not in allowed_expected:
                issues.append(
                    f"table {table_idx + 1} row {r} unexpected expected-result text: {row.cells[3].text}"
                )

    headers = [cell.text.strip() for cell in doc.tables[7].rows[0].cells]
    for required in ["password 模式", "challenge 模式", "otp(hotp) 模式", "mac 模式", "signature 模式"]:
        if required not in headers:
            issues.append(f"cross table missing header {required}")

    expected_detail_header = ["测试编号", "测试场景", "输入条件", "期望结果", "实际结果", "是否通过"]
    for table_idx in range(1, 7):
        actual_header = [cell.text.strip() for cell in doc.tables[table_idx].rows[0].cells]
        if actual_header != expected_detail_header:
            issues.append(f"table {table_idx + 1} header mismatch: {actual_header}")

    print(f"DOCX: {DOCX}")
    print(f"QA issues: {len(issues)}")
    for issue in issues:
        print(f"- {issue}")
    print(f"tables: {len(doc.tables)}")
    print(f"rows: {[len(t.rows) for t in doc.tables]}")
    print(f"cols: {[len(t.columns) for t in doc.tables]}")


if __name__ == "__main__":
    main()

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "output" / "experiment_results.json"
DOCX_OUT = ROOT / "实验五-第七部分实验结果表格.docx"


GROUPS = [
    ("1. 基线实验：静态口令鉴别", "PW-"),
    ("2. 挑战-应答鉴别实验", "CH-"),
    ("3.1 一次性口令 OTP：HOTP 计数器同步型", "HO-"),
    ("3.2 一次性口令 OTP：TOTP 时间同步型", "TO-"),
    ("4. MAC / HMAC 消息鉴别实验", "MC-"),
    ("5. 数字签名鉴别实验", "SG-"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def style_table(table, widths_cm: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, widths_cm)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "D9EAF7")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)
                run.font.name = "Microsoft YaHei"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, 70, 70, 70, 70)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(7.2)
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9)
    for name in ["Heading 1", "Heading 2"]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.color.rgb = RGBColor(31, 78, 121)


def make_actual_text(record: dict) -> str:
    actual = record["actual"]
    log = actual["log"]
    return (
        f"(a) 返回值/错误码：{actual['result']}, \"{actual['reason']}\"\n"
        f"(b) 服务器状态变化：{actual['state_before']}；{actual['state_after']}\n"
        f"(c) 关键日志字段：log.reason=\"{log['reason']}\", "
        f"log.result=\"{log['result']}\", log.server_time={log.get('server_time', '')}"
    )


def add_paragraph(doc: Document, text: str, bold=False, size=9, color=None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_reason_table(doc: Document) -> None:
    add_paragraph(doc, "拒绝原因类别说明", bold=True, size=10, color=(31, 78, 121))
    rows = [
        ("比较失败", "认证值不匹配", "口令错误、proof/tag/签名/OTP 计算错误、密钥不一致"),
        ("新鲜性失败", "请求超出有效时间或计数窗口", "nonce 过期、ts 过期、OTP 超出时间窗口、counter 漂移过远"),
        ("复用失败", "一次性凭证被重复提交", "nonce 已消费、同一 OTP 再次提交、同一时间片重复使用"),
        ("身份失败", "身份标识相关检查不通过", "用户不存在、账户已停用"),
        ("完整性失败", "消息体或元数据被篡改", "payload/id/ts 被修改导致 tag/签名不匹配"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "拒绝原因类别"
    table.rows[0].cells[1].text = "含义"
    table.rows[0].cells[2].text = "典型触发条件"
    for item in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(item):
            cells[idx].text = value
    style_table(table, [3.0, 5.0, 16.0])


def add_detail_group(doc: Document, title: str, records: list[dict]) -> None:
    doc.add_heading(title, level=2)
    headers = ["测试编号", "测试场景", "输入条件", "期望结果", "实际结果", "是否通过"]
    widths = [1.7, 3.5, 5.1, 2.7, 8.5, 5.0]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for rec in records:
        cells = table.add_row().cells
        values = [
            rec["test_id"],
            rec["scene"],
            rec["input"],
            rec["expected"],
            make_actual_text(rec),
            f"{rec['passed']}。{rec['analysis']}",
        ]
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, widths)


def add_cross_table(doc: Document, rows: list[dict]) -> None:
    doc.add_heading("6. 组合鉴别扩展实验：跨机制对比", level=2)
    headers = ["测试编号", "攻击类型", "password 模式", "challenge 模式", "otp(hotp) 模式", "mac 模式", "signature 模式"]
    widths = [1.7, 3.0, 4.1, 4.4, 4.3, 4.2, 4.5]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for rec in rows:
        cells = table.add_row().cells
        values = [rec["test_id"], rec["attack"], rec["password"], rec["challenge"], rec["otp(hotp)"], rec["mac"], rec["signature"]]
        for i, value in enumerate(values):
            cells[i].text = value
    style_table(table, widths)


def main() -> None:
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    records = data["records"]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)
    set_doc_styles(doc)

    title = doc.add_heading("七、实验结果", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_paragraph(
        doc,
        "本文件单独整理第七部分表格。实际结果由 main.py 运行生成，包含返回值/错误码、服务器状态变化、关键审计日志字段，以及期望与实际是否一致的判定说明。",
        size=9,
    )
    add_reason_table(doc)

    for title, prefix in GROUPS:
        group_records = [r for r in records if r["test_id"].startswith(prefix)]
        add_detail_group(doc, title, group_records)

    add_cross_table(doc, data["cross"])
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    main()
